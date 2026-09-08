from __future__ import annotations

import finess_sirene_review_export_engine as review_export
import merge_reviews
import pandas as pd
from openpyxl import load_workbook


def _canonical(value):
    if pd.isna(value):
        return None
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    return text or None


def _fill_review_workbook(path, decision):
    entity, siren, comment = decision
    workbook = load_workbook(path)
    worksheet = workbook["Review"] if "Review" in workbook.sheetnames else workbook.active
    headers = {
        str(cell.value): index
        for index, cell in enumerate(worksheet[1], start=1)
        if cell.value is not None
    }

    written = False
    for row_index in range(2, worksheet.max_row + 1):
        current = _canonical(
            worksheet.cell(row=row_index, column=headers["EJ"]).value
        )
        if current != entity or written:
            continue
        worksheet.cell(
            row=row_index,
            column=headers["Siren_retenu"],
            value=siren,
        )
        worksheet.cell(
            row=row_index,
            column=headers["Incertitude1"],
            value="0",
        )
        worksheet.cell(
            row=row_index,
            column=headers["Commentaire"],
            value=comment,
        )
        written = True

    assert written
    workbook.save(path)


def test_stage03_export_to_stage04_merge_with_two_reviewers(tmp_path, ej_review_config):
    proposals = pd.DataFrame(
        [
            {
                "EJ": "000000001",
                "candidate_rank": 1,
                "siren_proposal": "111111111",
                "stratum_id": "000001",
                "random_key": 0.1,
            },
            {
                "EJ": "000000001",
                "candidate_rank": 2,
                "siren_proposal": "121212121",
                "stratum_id": "000001",
                "random_key": 0.1,
            },
            {
                "EJ": "000000002",
                "candidate_rank": 1,
                "siren_proposal": "222222222",
                "stratum_id": "000002",
                "random_key": 0.2,
            },
            {
                "EJ": "000000002",
                "candidate_rank": 2,
                "siren_proposal": "232323232",
                "stratum_id": "000002",
                "random_key": 0.2,
            },
        ]
    )
    review_plan = pd.DataFrame(
        [
            {"stratum_id": "000001", "reviewer": "alice", "sample_size": 1},
            {"stratum_id": "000002", "reviewer": "bob", "sample_size": 1},
        ]
    )
    stage04_plan = review_plan.copy()
    review_root = tmp_path / "reviewed"

    review_export.export_review_files(
        proposals,
        review_plan,
        review_root,
        ej_review_config,
    )

    alice = review_root / "alice" / "000001__1.xlsx"
    bob = review_root / "bob" / "000002__1.xlsx"
    _fill_review_workbook(
        alice,
        ("000000001", "111111111", "reviewed by alice"),
    )
    _fill_review_workbook(
        bob,
        ("000000002", "222222222", "reviewed by bob"),
    )

    proposal_path = tmp_path / "ej_proposals.xlsx"
    proposals.to_excel(proposal_path, index=False)
    plan_path = tmp_path / "ej_strata_plan.xlsx"
    stage04_plan.to_excel(plan_path, index=False)
    correction_report = tmp_path / "corrections.csv"

    result = merge_reviews.run_merge_pipeline(
        results_root=review_root,
        config=merge_reviews.EJ_SIREN_CONFIG,
        proposal_path=proposal_path,
        plan_path=plan_path,
        correction_report_path=correction_report,
    )

    assert {loaded.path.name for loaded in result["loaded_files"]} == {
        "000001__1.xlsx",
        "000002__1.xlsx",
    }
    assert result["missing_strata"] == []
    assert pd.read_csv(correction_report).empty

    confirmed = result["confirmed_decisions_df"]
    assert dict(
        zip(
            confirmed["EJ"].map(_canonical),
            confirmed["Siren_retenu"].map(_canonical),
        )
    ) == {
        "000000001": "111111111",
        "000000002": "222222222",
    }
