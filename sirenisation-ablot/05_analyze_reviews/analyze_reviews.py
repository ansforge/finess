"""Per-stratum review analysis and strata-plan enrichment.

The enriched strata plan is the main analysis output. Global reporting and
threshold deployment scenarios are intentionally implemented separately in
``review_global_summary.py``.
"""

from __future__ import annotations

import math
import re
import textwrap
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from pathlib import Path

import finite_population_wilson_ci as finite_population_wilson
import pandas as pd
import sampling_precision

ENTITY_ID = "entity_id"
BUSINESS_ID = "business_id"
BUSINESS_ID_PROPOSAL = "business_id_proposal"
PROPOSAL_SOURCE = "proposal_source"
STRATUM_ID = "stratum_id"
UNCERTAINTY = "uncertainty"
ADRIEN_STATUS = "adrien_status"
IS_INITIAL_PROPOSAL = "is_initial_proposal"

_VALID_EXCEL_SUFFIXES = {".xlsx", ".xlsm", ".xltx", ".xltm"}


class DataValidationError(ValueError):
    """Raised when an analysis input violates the expected schema or rules."""


@dataclass(frozen=True)
class ReviewAnalysisConfig:
    """Map public dataframe columns and sampling-plan fields to the analysis model."""

    entity_name: str
    business_name: str
    entity_id_column: str
    business_id_proposal_column: str
    reviewed_business_id_column: str

    proposal_source_column: str = "proposal_source"
    stratum_id_column: str = "stratum_id"
    uncertainty_column: str = "Incertitude1"
    adrien_status_column: str = "statut_validation_Adrien"
    initial_proposal_flag_column: str | None = None

    # Target-entity population/sample fields in the strata plan.
    population_size_column: str | None = None
    sample_size_column: str | None = None
    sampling_weight_column: str | None = None

    # Legacy optional two-stage EGE design fields for workflows that already store
    # realized stage counts in the EGE strata plan. The current workflow instead
    # derives these counts from the EJ plan/sample and pooled EGE sampling frame and
    # passes them to ``enrich_strata_plan`` as ``ege_sampling_design_df``.
    first_stage_population_size_column: str | None = None
    first_stage_sample_size_column: str | None = None
    second_stage_pool_size_column: str | None = None
    second_stage_sample_size_column: str | None = None

    valid_business_id_specials: frozenset[str] = field(
        default_factory=lambda: frozenset({"NA", "N/A"})
    )
    valid_uncertainty_values: frozenset[str] = field(
        default_factory=lambda: frozenset({"-12", "-2", "-1", "0", "1", "2", "12"})
    )

    def __post_init__(self) -> None:
        public_columns = [
            self.entity_id_column,
            self.business_id_proposal_column,
            self.reviewed_business_id_column,
            self.proposal_source_column,
            self.stratum_id_column,
            self.uncertainty_column,
            self.adrien_status_column,
            self.effective_initial_proposal_flag_column,
        ]
        duplicates = sorted(
            {name for name in public_columns if public_columns.count(name) > 1}
        )
        if duplicates:
            raise ValueError(
                f"Configured dataframe column names must be unique: {duplicates}"
            )

    @property
    def effective_initial_proposal_flag_column(self) -> str:
        if self.initial_proposal_flag_column:
            return self.initial_proposal_flag_column
        token = re.sub(
            r"[^0-9A-Za-z]+", "_", str(self.business_name).strip()
        ).strip("_").lower()
        if not token:
            raise ValueError(
                f"Cannot derive an initial-proposal flag from {self.business_name!r}"
            )
        return f"is_initial_{token}"

    @property
    def effective_population_size_column(self) -> str:
        if self.population_size_column:
            return self.population_size_column
        return f"n_{self.entity_name.upper()}"

    @property
    def public_to_internal(self) -> Mapping[str, str]:
        return {
            self.entity_id_column: ENTITY_ID,
            self.business_id_proposal_column: BUSINESS_ID_PROPOSAL,
            self.reviewed_business_id_column: BUSINESS_ID,
            self.proposal_source_column: PROPOSAL_SOURCE,
            self.stratum_id_column: STRATUM_ID,
            self.uncertainty_column: UNCERTAINTY,
            self.adrien_status_column: ADRIEN_STATUS,
            self.effective_initial_proposal_flag_column: IS_INITIAL_PROPOSAL,
        }

    @property
    def internal_to_public(self) -> Mapping[str, str]:
        return {internal: public for public, internal in self.public_to_internal.items()}

    @property
    def entity_plural(self) -> str:
        return f"{self.entity_name.upper()}s"

    @property
    def business_upper(self) -> str:
        return self.business_name.upper()

    @property
    def is_ege(self) -> bool:
        return self.entity_name.upper() == "EGE"

    @property
    def analysis_columns(self) -> dict[str, str]:
        """Map semantic analysis names to human-readable output-column names."""

        e = self.entity_plural
        b = self.business_upper
        return {
            "stratum_id": "Stratum ID",
            "population_size": "Population size",
            "sample_size": "Sample size",  # internal design label; not displayed
            "sampling_weight": "Sampling weight",
            # Extra design fields shown only for EGE output.
            "first_stage_population_size": "EJ population size",
            "first_stage_sample_size": "Sampled EJs",
            "second_stage_pool_size": "Pooled EGEs in sampled EJs",
            "second_stage_sample_size": "Sampled EGEs",
            "first_stage_inclusion_probability": "EJ inclusion probability",
            "second_stage_inclusion_probability": "EGE conditional inclusion probability",
            # Review counts.
            "reviewed": f"{e} examined",
            "sampling_moe_95": "Sampling margin of error — 95%",
            "to_close": f"{e} to close",
            "to_match": f"{e} to match",
            "initial_correct": "Initial correct",
            "initial_na": "Initial NA",
            "initial_na_correct": "Initial NA correct",
            "ans_proposed": "ANS proposal available",
            "adrien_only_proposed": "Adrien — only proposal rule available",
            "adrien_proposed": "Adrien proposal available",
            "ans_correct": "ANS correct when proposed",
            "adrien_only_correct": "Adrien correct — only proposal, when applicable",
            "adrien_correct": "Adrien correct when proposed",
            "ans_fallback_correct": "ANS + Initial fallback correct",
            "adrien_only_fallback_correct": "Adrien — only proposal + Initial fallback correct",
            "adrien_fallback_correct": "Adrien + Initial fallback correct",
            "ans_fallback_mistakes_corrected": "ANS + Initial fallback — mistakes corrected",
            "ans_fallback_mistakes_introduced": "ANS + Initial fallback — mistakes introduced",
            "adrien_fallback_mistakes_corrected": "Adrien + Initial fallback — mistakes corrected",
            "adrien_fallback_mistakes_introduced": "Adrien + Initial fallback — mistakes introduced",
            "adrien_only_fallback_mistakes_corrected": "Adrien — only proposal + Initial fallback — mistakes corrected",
            "adrien_only_fallback_mistakes_introduced": "Adrien — only proposal + Initial fallback — mistakes introduced",
            "neither_correct": "Neither ANS nor Adrien correct",
            "with_uncertainty": f"{e} with uncertainty",
            "business_identified": f"{e} with {b} identified",
            # Review proportions.
            "initial_correct_pct": "Initial correct (%)",
            "initial_na_pct": "Initial NA (%)",
            "initial_na_correct_pct": "Initial NA correct (%)",
            "ans_coverage_pct": "ANS coverage (%)",
            "adrien_only_coverage_pct": "Adrien — only proposal coverage (%)",
            "adrien_coverage_pct": "Adrien coverage (%)",
            "ans_correct_pct": "ANS correct when proposed (%)",
            "adrien_only_correct_pct": "Adrien correct — only proposal, when applicable (%)",
            "adrien_correct_pct": "Adrien correct when proposed (%)",
            "ans_fallback_correct_pct": "ANS + Initial fallback correct (%)",
            "adrien_only_fallback_correct_pct": "Adrien — only proposal + Initial fallback correct (%)",
            "adrien_fallback_correct_pct": "Adrien + Initial fallback correct (%)",
            "neither_correct_pct": "Neither ANS nor Adrien correct (%)",
            "with_uncertainty_pct": f"{e} with uncertainty (%)",
            "to_close_pct": f"{e} to close (%)",
            "business_identified_pct": f"{e} with {b} identified (%)",
            # Wilson intervals. These are descriptive only.
            "initial_wilson_lower": "Initial correct — Wilson 95% lower",
            "initial_wilson_upper": "Initial correct — Wilson 95% upper",
            "ans_wilson_lower": "ANS correct when proposed — Wilson 95% lower",
            "ans_wilson_upper": "ANS correct when proposed — Wilson 95% upper",
            "adrien_wilson_lower": "Adrien correct when proposed — Wilson 95% lower",
            "adrien_wilson_upper": "Adrien correct when proposed — Wilson 95% upper",
            # Accepted/satisfactory classifications use deployable rules: use the
            # method when it is available for an entity, otherwise keep Initial.
            "initial_ge_95": "Initial correctness ≥ 95%",
            "initial_ge_99": "Initial correctness ≥ 99%",
            "ans_ge_95": "ANS + Initial fallback correctness ≥ 95%",
            "ans_ge_99": "ANS + Initial fallback correctness ≥ 99%",
            "adrien_ge_95": "Adrien + Initial fallback correctness ≥ 95%",
            "adrien_ge_99": "Adrien + Initial fallback correctness ≥ 99%",
            "adrien_only_ge_95": "Adrien — only proposal + Initial fallback correctness ≥ 95%",
            "adrien_only_ge_99": "Adrien — only proposal + Initial fallback correctness ≥ 99%",
        }

    @property
    def requested_output_order(self) -> list[str]:
        names = self.analysis_columns
        design = [
            names["stratum_id"],
            names["population_size"],
            names["reviewed"],
            names["sampling_moe_95"],
            names["sampling_weight"],
        ]
        if self.is_ege:
            design.extend(
                [
                    names["first_stage_population_size"],
                    names["first_stage_sample_size"],
                    names["second_stage_pool_size"],
                    names["second_stage_sample_size"],
                    names["first_stage_inclusion_probability"],
                    names["second_stage_inclusion_probability"],
                ]
            )
        return design + [
            names["to_close"],
            names["to_match"],
            names["initial_correct"],
            names["initial_na"],
            names["initial_na_correct"],
            names["ans_proposed"],
            names["adrien_only_proposed"],
            names["adrien_proposed"],
            names["ans_correct"],
            names["adrien_only_correct"],
            names["adrien_correct"],
            names["ans_fallback_correct"],
            names["adrien_only_fallback_correct"],
            names["adrien_fallback_correct"],
            names["ans_fallback_mistakes_corrected"],
            names["ans_fallback_mistakes_introduced"],
            names["adrien_fallback_mistakes_corrected"],
            names["adrien_fallback_mistakes_introduced"],
            names["adrien_only_fallback_mistakes_corrected"],
            names["adrien_only_fallback_mistakes_introduced"],
            names["neither_correct"],
            names["with_uncertainty"],
            names["business_identified"],
            names["initial_correct_pct"],
            names["initial_na_pct"],
            names["initial_na_correct_pct"],
            names["ans_coverage_pct"],
            names["adrien_only_coverage_pct"],
            names["adrien_coverage_pct"],
            names["ans_correct_pct"],
            names["adrien_only_correct_pct"],
            names["adrien_correct_pct"],
            names["ans_fallback_correct_pct"],
            names["adrien_only_fallback_correct_pct"],
            names["adrien_fallback_correct_pct"],
            names["neither_correct_pct"],
            names["with_uncertainty_pct"],
            names["to_close_pct"],
            names["business_identified_pct"],
            names["initial_wilson_lower"],
            names["initial_wilson_upper"],
            names["ans_wilson_lower"],
            names["ans_wilson_upper"],
            names["adrien_wilson_lower"],
            names["adrien_wilson_upper"],
            names["initial_ge_95"],
            names["initial_ge_99"],
            names["ans_ge_95"],
            names["ans_ge_99"],
            names["adrien_ge_95"],
            names["adrien_ge_99"],
            names["adrien_only_ge_95"],
            names["adrien_only_ge_99"],
        ]


EJ_SIREN_CONFIG = ReviewAnalysisConfig(
    entity_name="EJ",
    business_name="siren",
    entity_id_column="EJ",
    business_id_proposal_column="siren_proposal",
    reviewed_business_id_column="Siren_retenu",
)

EGE_SIRET_CONFIG = ReviewAnalysisConfig(
    entity_name="EGE",
    business_name="siret",
    entity_id_column="EGE",
    business_id_proposal_column="siret_proposal",
    reviewed_business_id_column="Siret_retenu",
    adrien_status_column="statut_validation_Adrien_EGE",
)


def _canonical_text(value) -> str | None:
    if pd.isna(value):
        return None
    text = str(value).strip()
    return text if text else None


def _canonical_identifier(value) -> str | None:
    return _canonical_text(value)


def _is_true(value) -> bool:
    if isinstance(value, bool):
        return value
    if pd.isna(value):
        return False
    return str(value).strip().lower() in {"1", "true", "yes", "y"}


def _business_is_na(value: str | None, config: ReviewAnalysisConfig) -> bool:
    return value is None or value in config.valid_business_id_specials


def _raise_if_missing_columns(
    df: pd.DataFrame, columns: Sequence[str], *, source: str
) -> None:
    missing = [column for column in columns if column not in df.columns]
    if missing:
        raise DataValidationError(f"{source}: missing required columns: {missing}")


def _to_internal(
    df: pd.DataFrame,
    config: ReviewAnalysisConfig,
    *,
    required_internal: Sequence[str] = (),
    source: str = "dataframe",
) -> pd.DataFrame:
    required_public = [config.internal_to_public[name] for name in required_internal]
    _raise_if_missing_columns(df, required_public, source=source)

    rename_map = {
        public: internal
        for public, internal in config.public_to_internal.items()
        if public in df.columns and public != internal
    }
    collisions = sorted(
        internal
        for public, internal in rename_map.items()
        if internal in df.columns and internal != public
    )
    if collisions:
        raise DataValidationError(
            f"{source}: neutral internal column name collision(s): {collisions}"
        )
    return df.copy().rename(columns=rename_map)


def _read_dataframe(path: str | Path) -> pd.DataFrame:
    source = Path(path)
    if not source.exists():
        raise FileNotFoundError(f"Input file not found: {source}")

    suffix = source.suffix.lower()
    if suffix in _VALID_EXCEL_SUFFIXES:
        return pd.read_excel(source, dtype=str, keep_default_na=False)
    if suffix == ".parquet":
        return pd.read_parquet(source)
    raise DataValidationError(
        f"Unsupported input format for {source.name!r}; expected Excel or parquet"
    )


def load_confirmed_decisions(
    confirmed_decisions_path: str | Path,
    *,
    config: ReviewAnalysisConfig = EJ_SIREN_CONFIG,
) -> pd.DataFrame:
    """Load the one-row-per-entity confirmed-decisions database."""

    df = _read_dataframe(confirmed_decisions_path)
    _raise_if_missing_columns(
        df,
        (
            config.entity_id_column,
            config.reviewed_business_id_column,
            config.uncertainty_column,
        ),
        source=str(confirmed_decisions_path),
    )

    normalized_entities = df[config.entity_id_column].map(_canonical_identifier)
    if normalized_entities.isna().any():
        raise DataValidationError(
            f"{confirmed_decisions_path}: missing {config.entity_id_column} values"
        )
    duplicate_mask = normalized_entities.duplicated(keep=False)
    if duplicate_mask.any():
        duplicates = sorted(set(normalized_entities.loc[duplicate_mask]))
        raise DataValidationError(
            f"{confirmed_decisions_path}: confirmed decisions must contain one row per "
            f"{config.entity_name}; duplicates: {', '.join(duplicates)}"
        )
    return df


def load_proposal_table(
    proposal_path: str | Path,
    *,
    config: ReviewAnalysisConfig = EJ_SIREN_CONFIG,
) -> pd.DataFrame:
    """Load proposal rows required to analyze confirmed reviewer decisions."""

    df = _read_dataframe(proposal_path)
    _raise_if_missing_columns(
        df,
        (
            config.entity_id_column,
            config.business_id_proposal_column,
            config.proposal_source_column,
            config.stratum_id_column,
            config.adrien_status_column,
            config.effective_initial_proposal_flag_column,
        ),
        source=str(proposal_path),
    )
    return df


def load_strata_plan(
    plan_path: str | Path,
    *,
    config: ReviewAnalysisConfig = EJ_SIREN_CONFIG,
) -> pd.DataFrame:
    """Load the strata plan used as the target of the enrichment step."""

    df = _read_dataframe(plan_path)
    _raise_if_missing_columns(
        df,
        (config.stratum_id_column, config.effective_population_size_column),
        source=str(plan_path),
    )
    return df



def validate_ege_pooled_frame_against_plan(
    ege_plan_df: pd.DataFrame,
    pooled_ege_proposals_df: pd.DataFrame,
    *,
    ege_id_column: str = "EGE",
    stratum_id_column: str = "stratum_id",
    pooled_count_column: str = "n_EGE_pooled",
) -> dict[str, int]:
    """Validate the pooled EGE frame against the Stage 02 EGE strata plan.

    The EGE plan retains ``n_EGE_pooled`` as the number of unique pooled EGE in
    each inherited stratum. Stage 05 uses this check before reconstructing the
    realised two-stage sampling design so a stale or mismatched pooled frame
    cannot silently affect the analysis.
    """
    _raise_if_missing_columns(
        ege_plan_df,
        (stratum_id_column, pooled_count_column),
        source="ege_plan_df",
    )
    _raise_if_missing_columns(
        pooled_ege_proposals_df,
        (ege_id_column, stratum_id_column),
        source="pooled_ege_proposals_df",
    )

    pooled = pooled_ege_proposals_df[[ege_id_column, stratum_id_column]].copy()
    pooled["_ege_norm"] = pooled[ege_id_column].map(_canonical_identifier)
    pooled["_stratum_norm"] = pooled[stratum_id_column].map(_canonical_identifier)
    if pooled[["_ege_norm", "_stratum_norm"]].isna().any().any():
        raise DataValidationError(
            "pooled_ege_proposals_df contains missing EGE or stratum identifiers"
        )

    ege_strata = pooled[["_ege_norm", "_stratum_norm"]].drop_duplicates()
    ambiguous = ege_strata["_ege_norm"].duplicated(keep=False)
    if ambiguous.any():
        examples = sorted(
            set(ege_strata.loc[ambiguous, "_ege_norm"].tolist())
        )[:20]
        raise DataValidationError(
            "Each pooled EGE must belong to exactly one stratum; ambiguous EGE "
            "include: " + ", ".join(examples)
        )

    pooled_counts = (
        ege_strata.groupby("_stratum_norm", sort=False)["_ege_norm"]
        .nunique()
        .astype("int64")
    )

    plan = ege_plan_df[[stratum_id_column, pooled_count_column]].copy()
    plan["_stratum_norm"] = plan[stratum_id_column].map(_canonical_identifier)
    if plan["_stratum_norm"].isna().any():
        raise DataValidationError("ege_plan_df contains missing stratum identifiers")
    if plan["_stratum_norm"].duplicated().any():
        duplicates = sorted(
            set(plan.loc[plan["_stratum_norm"].duplicated(False), "_stratum_norm"])
        )
        raise DataValidationError(
            "ege_plan_df must contain one row per stratum; duplicates: "
            + ", ".join(duplicates[:20])
        )

    plan_counts = pd.to_numeric(plan[pooled_count_column], errors="coerce")
    invalid = plan_counts.isna() | plan_counts.lt(0) | plan_counts.mod(1).ne(0)
    if invalid.any():
        sample = plan.loc[
            invalid,
            [stratum_id_column, pooled_count_column],
        ].head(20).to_dict("records")
        raise DataValidationError(
            f"ege_plan_df[{pooled_count_column!r}] must contain non-negative "
            f"integer counts. Invalid values: {sample}"
        )
    plan["_pooled_from_plan"] = plan_counts.astype("int64")
    plan["_pooled_from_frame"] = (
        plan["_stratum_norm"].map(pooled_counts).fillna(0).astype("int64")
    )

    mismatch = plan["_pooled_from_plan"].ne(plan["_pooled_from_frame"])
    extra_strata = sorted(set(pooled_counts.index) - set(plan["_stratum_norm"]))
    if mismatch.any() or extra_strata:
        sample = plan.loc[
            mismatch,
            [stratum_id_column, "_pooled_from_plan", "_pooled_from_frame"],
        ].head(20).to_dict("records")
        raise DataValidationError(
            f"{pooled_count_column} does not match the pooled EGE frame. "
            f"Extra pooled strata: {extra_strata[:10]}. "
            f"Count mismatches: {sample}"
        )

    return {
        "pooled_ege_count": int(ege_strata["_ege_norm"].nunique()),
        "pooled_strata_count": int(pooled_counts.size),
        "plan_strata_count": len(plan),
    }


def build_ege_two_stage_sampling_design(
    ej_plan_df: pd.DataFrame,
    ej_confirmed_decisions_df: pd.DataFrame,
    pooled_ege_proposals_df: pd.DataFrame,
    *,
    ej_id_column: str = "EJ",
    ege_id_column: str = "EGE",
    stratum_id_column: str = "stratum_id",
    ej_population_size_column: str = "n_EJ",
) -> pd.DataFrame:
    """Derive the realized two-stage EGE sampling design by stratum.

    The first-stage sample is the set of EJs in the confirmed-EJ file. The
    pooled EGE proposal table is expected to have already been restricted to
    those EJs; it supplies both the EJ-to-stratum mapping and the second-stage
    EGE sampling frame. Proposal rows are deduplicated by EGE before the pooled
    frame size is counted.

    The returned dataframe contains the design counts needed by
    :func:`enrich_strata_plan`. The completed EGE review count is intentionally
    not calculated here; it is taken from the per-stratum EGE analysis so the
    design is checked against the actual reviewed sample.
    """

    _raise_if_missing_columns(
        ej_plan_df,
        (stratum_id_column, ej_population_size_column),
        source="ej_plan_df",
    )
    _raise_if_missing_columns(
        ej_confirmed_decisions_df, (ej_id_column,), source="ej_confirmed_decisions_df"
    )
    _raise_if_missing_columns(
        pooled_ege_proposals_df,
        (ej_id_column, ege_id_column, stratum_id_column),
        source="pooled_ege_proposals_df",
    )

    selected = ej_confirmed_decisions_df[[ej_id_column]].copy()
    selected["_ej_norm"] = selected[ej_id_column].map(_canonical_identifier)
    if selected["_ej_norm"].isna().any():
        raise DataValidationError("ej_confirmed_decisions_df: missing EJ identifiers")
    if selected["_ej_norm"].duplicated().any():
        duplicates = sorted(
            set(selected.loc[selected["_ej_norm"].duplicated(keep=False), "_ej_norm"])
        )
        raise DataValidationError(
            "ej_confirmed_decisions_df must contain one row per EJ; duplicates: "
            + ", ".join(duplicates)
        )

    pooled = pooled_ege_proposals_df[
        [ej_id_column, ege_id_column, stratum_id_column]
    ].copy()
    pooled["_ej_norm"] = pooled[ej_id_column].map(_canonical_identifier)
    pooled["_ege_norm"] = pooled[ege_id_column].map(_canonical_identifier)
    pooled["_stratum_norm"] = pooled[stratum_id_column].map(_canonical_identifier)
    if pooled[["_ej_norm", "_ege_norm", "_stratum_norm"]].isna().any().any():
        raise DataValidationError(
            "pooled_ege_proposals_df must have non-missing EJ, EGE and stratum identifiers"
        )

    selected_ids = set(selected["_ej_norm"])
    pooled_parent_ids = set(pooled["_ej_norm"])
    unexpected = sorted(pooled_parent_ids - selected_ids)
    if unexpected:
        preview = ", ".join(unexpected[:10])
        suffix = " ..." if len(unexpected) > 10 else ""
        raise DataValidationError(
            "pooled_ege_proposals_df contains parent EJs that are absent from the "
            f"confirmed-EJ sample: {preview}{suffix}"
        )

    # Every selected EJ should appear in the pooled EGE frame because EJs with
    # no EGE are outside the EJ analysis population. Requiring this also avoids
    # silently understating the first-stage sample count.
    missing_selected = sorted(selected_ids - pooled_parent_ids)
    if missing_selected:
        preview = ", ".join(missing_selected[:10])
        suffix = " ..." if len(missing_selected) > 10 else ""
        raise DataValidationError(
            "Some confirmed EJs have no row in pooled_ege_proposals_df, so their "
            "stratum cannot be recovered for the EGE sampling design: "
            f"{preview}{suffix}"
        )

    ej_strata = (
        pooled[["_ej_norm", "_stratum_norm"]]
        .drop_duplicates()
        .groupby("_ej_norm", sort=False)["_stratum_norm"]
        .agg(list)
    )
    ambiguous_ejs = {
        ej: sorted(set(strata))
        for ej, strata in ej_strata.items()
        if len(set(strata)) != 1
    }
    if ambiguous_ejs:
        details = "; ".join(
            f"{ej}: {', '.join(strata)}"
            for ej, strata in list(ambiguous_ejs.items())[:10]
        )
        raise DataValidationError(
            "Selected EJs must map to exactly one stratum in the pooled EGE frame: "
            + details
        )

    selected_map = selected.copy()
    selected_map["_stratum_norm"] = selected_map["_ej_norm"].map(
        {ej: strata[0] for ej, strata in ej_strata.items()}
    )
    selected_counts = (
        selected_map.groupby("_stratum_norm", sort=False)["_ej_norm"]
        .nunique()
        .rename("Sampled EJs")
    )

    # One EGE can have several proposal rows, but it is one unit in the sampling
    # frame. Also require an EGE to belong to one stratum only.
    ege_strata = pooled[["_ege_norm", "_stratum_norm"]].drop_duplicates()
    ambiguous_eges = ege_strata["_ege_norm"].duplicated(keep=False)
    if ambiguous_eges.any():
        examples = sorted(
            set(ege_strata.loc[ambiguous_eges, "_ege_norm"].tolist())
        )[:10]
        raise DataValidationError(
            "Pooled EGEs must map to exactly one stratum; ambiguous EGEs include: "
            + ", ".join(examples)
        )
    pooled_counts = (
        ege_strata.groupby("_stratum_norm", sort=False)["_ege_norm"]
        .nunique()
        .rename("Pooled EGEs in sampled EJs")
    )

    ej_plan = ej_plan_df[[stratum_id_column, ej_population_size_column]].copy()
    ej_plan["_stratum_norm"] = ej_plan[stratum_id_column].map(_canonical_identifier)
    if ej_plan["_stratum_norm"].isna().any():
        raise DataValidationError("ej_plan_df: missing stratum identifiers")
    if ej_plan["_stratum_norm"].duplicated().any():
        raise DataValidationError("ej_plan_df must contain one row per stratum")
    ej_population = pd.to_numeric(ej_plan[ej_population_size_column], errors="coerce")
    if ej_population.isna().any() or (ej_population <= 0).any():
        raise DataValidationError(
            f"ej_plan_df[{ej_population_size_column!r}] must be positive and numeric"
        )
    ej_plan["EJ population size"] = ej_population

    design = (
        ej_plan[["_stratum_norm", "EJ population size"]]
        .merge(selected_counts, left_on="_stratum_norm", right_index=True, how="left")
        .merge(pooled_counts, left_on="_stratum_norm", right_index=True, how="left")
    )
    design[["Sampled EJs", "Pooled EGEs in sampled EJs"]] = design[
        ["Sampled EJs", "Pooled EGEs in sampled EJs"]
    ].fillna(0)

    # Strata not reached by the realized EJ sample do not contribute reviewed EGEs
    # and therefore do not need a two-stage weight. Keep only sampled strata.
    design = design.loc[design["Sampled EJs"].gt(0)].copy()
    if (design["Pooled EGEs in sampled EJs"] <= 0).any():
        bad = design.loc[
            design["Pooled EGEs in sampled EJs"] <= 0, "_stratum_norm"
        ].tolist()
        raise DataValidationError(
            "Sampled EJ strata must have a non-empty pooled EGE frame: "
            + ", ".join(map(str, bad))
        )
    if (design["Sampled EJs"] > design["EJ population size"]).any():
        raise DataValidationError(
            "The number of sampled EJs cannot exceed the EJ population in a stratum."
        )

    design["Stratum ID"] = design["_stratum_norm"]
    for column in [
        "EJ population size",
        "Sampled EJs",
        "Pooled EGEs in sampled EJs",
    ]:
        design[column] = pd.to_numeric(design[column], errors="raise").round().astype("Int64")

    return design[
        [
            "Stratum ID",
            "EJ population size",
            "Sampled EJs",
            "Pooled EGEs in sampled EJs",
        ]
    ].reset_index(drop=True)


def _uncertainty_is_present(value: str | None, config: ReviewAnalysisConfig) -> bool:
    value = _canonical_text(value)
    if value not in config.valid_uncertainty_values:
        return False
    try:
        return int(value) != 0
    except (TypeError, ValueError):
        return False


def _uncertainty_closes_structure(
    value: str | None, config: ReviewAnalysisConfig
) -> bool:
    value = _canonical_text(value)
    if value not in config.valid_uncertainty_values:
        return False
    try:
        return int(value) < 0
    except (TypeError, ValueError):
        return False


def _wilson_interval_pct(successes: int, trials: int) -> tuple[object, object]:
    """Return the ordinary 95% Wilson interval before plan enrichment.

    ``enrich_strata_plan`` overwrites these provisional bounds with the
    finite-population Wilson-type interval once the stratum population and
    realized examined count are available.
    """

    if trials <= 0:
        return pd.NA, pd.NA
    if successes < 0 or successes > trials:
        raise DataValidationError(
            f"Wilson interval requires 0 <= successes <= trials; got {successes}/{trials}"
        )

    z = 1.959963984540054
    p = successes / trials
    z2 = z * z
    denominator = 1 + z2 / trials
    center = (p + z2 / (2 * trials)) / denominator
    half = (
        z
        * math.sqrt(p * (1 - p) / trials + z2 / (4 * trials * trials))
        / denominator
    )
    return round(100 * max(0.0, center - half), 2), round(
        100 * min(1.0, center + half), 2
    )


def _pct(numerator: int, denominator: int) -> object:
    if denominator <= 0:
        return pd.NA
    return round(100 * numerator / denominator, 2)


def _finite_population_wilson_interval_pct(
    population_size: object,
    examined: object,
    trials: object,
    successes: object,
    *,
    stratum_id: object,
    metric_label: str,
) -> tuple[object, object]:
    """Return the finite-population Wilson-type 95% CI in percentage points.

    The observed proportion is ``successes / trials``. ``trials`` is the number
    of entities to match for Initial, and the number with a method proposal for
    conditional ANS/Adrien correctness. The finite-population correction uses the
    actual examined count versus total stratum population.
    """

    if pd.isna(examined) or float(examined) <= 0:
        return pd.NA, pd.NA
    if pd.isna(trials) or float(trials) <= 0:
        return pd.NA, pd.NA
    if pd.isna(successes):
        raise DataValidationError(
            f"Missing {metric_label} success count for reviewed stratum {stratum_id!r}."
        )

    try:
        lower, upper = finite_population_wilson.finite_population_wilson_ci(
            population_size,
            trials,
            successes,
            confidence=0.95,
            fpc_sample_size=examined,
        )
    except ValueError as exc:
        raise DataValidationError(
            f"Cannot compute finite-population Wilson interval for {metric_label} "
            f"in stratum {stratum_id!r}: {exc}"
        ) from exc

    if math.isnan(lower) or math.isnan(upper):
        return pd.NA, pd.NA
    return round(100.0 * lower, 2), round(100.0 * upper, 2)


def build_entity_review_outcomes(
    confirmed_decisions_df: pd.DataFrame,
    proposal_df: pd.DataFrame,
    *,
    config: ReviewAnalysisConfig = EJ_SIREN_CONFIG,
) -> pd.DataFrame:
    """Build one row of review-outcome indicators per reviewed entity.

    Entities identified for closure remain in overall review statistics but are
    excluded from proposal-performance counts. Initial and deployable fallback
    correctness use entities to match; ANS/Adrien conditional correctness uses
    only entities where the corresponding method is available.
    """

    decisions = _to_internal(
        confirmed_decisions_df,
        config,
        required_internal=(ENTITY_ID, BUSINESS_ID, UNCERTAINTY),
        source="confirmed_decisions_df",
    )
    proposals = _to_internal(
        proposal_df,
        config,
        required_internal=(
            ENTITY_ID,
            BUSINESS_ID_PROPOSAL,
            PROPOSAL_SOURCE,
            STRATUM_ID,
            ADRIEN_STATUS,
            IS_INITIAL_PROPOSAL,
        ),
        source="proposal_df",
    )

    decisions["_entity_id_norm"] = decisions[ENTITY_ID].map(_canonical_identifier)
    proposals["_entity_id_norm"] = proposals[ENTITY_ID].map(_canonical_identifier)
    proposals["_stratum_id_norm"] = proposals[STRATUM_ID].map(_canonical_identifier)
    proposals["_business_id_proposal_norm"] = proposals[BUSINESS_ID_PROPOSAL].map(
        _canonical_identifier
    )
    proposals["_is_initial_proposal_norm"] = proposals[IS_INITIAL_PROPOSAL].map(_is_true)

    if decisions["_entity_id_norm"].isna().any():
        raise DataValidationError("confirmed_decisions_df: missing entity identifiers")

    duplicate_decisions = decisions["_entity_id_norm"].duplicated(keep=False)
    if duplicate_decisions.any():
        duplicates = sorted(set(decisions.loc[duplicate_decisions, "_entity_id_norm"]))
        raise DataValidationError(
            "confirmed_decisions_df must contain one row per "
            f"{config.entity_name}; duplicates: {', '.join(duplicates)}"
        )

    proposal_strata: dict[str, set[str]] = {}
    for entity_id, group in proposals.groupby("_entity_id_norm", dropna=False, sort=False):
        if entity_id is None:
            continue
        proposal_strata[entity_id] = {
            value for value in group["_stratum_id_norm"].tolist() if value is not None
        }

    decision_entities = decisions["_entity_id_norm"].tolist()
    decision_entity_set = set(decision_entities)
    missing_entities = sorted(
        entity_id for entity_id in decision_entities if entity_id not in proposal_strata
    )
    if missing_entities:
        raise DataValidationError(
            f"Confirmed {config.entity_name}s absent from proposal_df: "
            + ", ".join(missing_entities)
        )

    ambiguous_entities = {
        entity_id: strata
        for entity_id, strata in proposal_strata.items()
        if entity_id in decision_entity_set and len(strata) != 1
    }
    if ambiguous_entities:
        details = "; ".join(
            f"{entity_id}: {', '.join(sorted(strata)) if strata else '<missing stratum>'}"
            for entity_id, strata in sorted(ambiguous_entities.items())
        )
        raise DataValidationError(
            f"Confirmed {config.entity_name}s without exactly one proposal stratum: {details}"
        )

    initial_counts = (
        proposals.loc[proposals["_entity_id_norm"].isin(decision_entity_set)]
        .groupby("_entity_id_norm")["_is_initial_proposal_norm"]
        .sum()
    )
    invalid_initial_rows = {
        entity_id: int(initial_counts.get(entity_id, 0))
        for entity_id in decision_entities
        if int(initial_counts.get(entity_id, 0)) != 1
    }
    if invalid_initial_rows:
        details = "; ".join(
            f"{entity_id}: {count}"
            for entity_id, count in sorted(invalid_initial_rows.items())
        )
        raise DataValidationError(
            f"Confirmed {config.entity_name}s must have exactly one initial proposal row "
            f"({config.effective_initial_proposal_flag_column}=True) in proposal_df. "
            f"Found counts: {details}"
        )

    rows: list[dict[str, object]] = []
    for _, decision_row in decisions.iterrows():
        entity_id = decision_row["_entity_id_norm"]
        stratum_id = next(iter(proposal_strata[entity_id]))
        selected_business = _canonical_identifier(decision_row[BUSINESS_ID])
        selected_uncertainty = _canonical_text(decision_row[UNCERTAINTY])
        entity_proposals = proposals.loc[proposals["_entity_id_norm"].eq(entity_id)]
        initial_row = entity_proposals.loc[
            entity_proposals["_is_initial_proposal_norm"]
        ].iloc[0]
        initial_business = _canonical_identifier(initial_row[BUSINESS_ID_PROPOSAL])

        if _business_is_na(selected_business, config):
            candidate_matches = entity_proposals.iloc[0:0]
        else:
            candidate_matches = entity_proposals.loc[
                entity_proposals["_business_id_proposal_norm"].eq(selected_business)
            ]

        closes_structure = _uncertainty_closes_structure(
            selected_uncertainty, config
        )
        to_match = not closes_structure
        initial_is_na = _business_is_na(initial_business, config)
        selected_is_na = _business_is_na(selected_business, config)

        initial_na = int(to_match and initial_is_na)
        initial_na_correct = int(to_match and initial_is_na and selected_is_na)
        initial_correct = int(
            to_match
            and (
                (initial_is_na and selected_is_na)
                or (
                    not initial_is_na
                    and not selected_is_na
                    and initial_business == selected_business
                )
            )
        )

        # Availability is entity-level and independent of whether the proposal
        # matches the reviewed business. This distinction is essential for the
        # deployable fallback rules: use a method when it proposed something,
        # otherwise retain Initial.
        entity_sources = {
            source
            for value in entity_proposals[PROPOSAL_SOURCE].tolist()
            if (source := _canonical_text(value)) is not None
        }
        matched_sources = {
            source
            for value in candidate_matches[PROPOSAL_SOURCE].tolist()
            if (source := _canonical_text(value)) is not None
        }

        ans_proposed = int(to_match and bool(entity_sources & {"ANS", "both"}))
        adrien_proposed = int(to_match and bool(entity_sources & {"Adrien", "both"}))
        adrien_only_proposed = int(
            to_match
            and adrien_proposed
            and any(
                _canonical_text(row[PROPOSAL_SOURCE]) in {"Adrien", "both"}
                and _canonical_text(row[ADRIEN_STATUS]) == "1 proposition"
                for _, row in entity_proposals.iterrows()
            )
        )

        ans_correct = int(to_match and bool(matched_sources & {"ANS", "both"}))
        adrien_correct = int(to_match and bool(matched_sources & {"Adrien", "both"}))
        adrien_only_correct = int(adrien_only_proposed and adrien_correct)

        ans_fallback_correct = int(
            to_match and (ans_correct or (not ans_proposed and initial_correct))
        )
        adrien_fallback_correct = int(
            to_match and (adrien_correct or (not adrien_proposed and initial_correct))
        )
        adrien_only_fallback_correct = int(
            to_match
            and (
                adrien_only_correct
                or (not adrien_only_proposed and initial_correct)
            )
        )

        # Joint Initial-versus-deployed-rule transitions. These cannot be
        # reconstructed from marginal correctness totals alone, so compute them
        # while the entity-level outcomes are still available. With fallback, a
        # missing alternative proposal never introduces a mistake because Initial
        # is retained in that case.
        ans_fallback_mistakes_corrected = int(
            to_match and not initial_correct and ans_fallback_correct
        )
        ans_fallback_mistakes_introduced = int(
            to_match and initial_correct and not ans_fallback_correct
        )
        adrien_fallback_mistakes_corrected = int(
            to_match and not initial_correct and adrien_fallback_correct
        )
        adrien_fallback_mistakes_introduced = int(
            to_match and initial_correct and not adrien_fallback_correct
        )
        adrien_only_fallback_mistakes_corrected = int(
            to_match and not initial_correct and adrien_only_fallback_correct
        )
        adrien_only_fallback_mistakes_introduced = int(
            to_match and initial_correct and not adrien_only_fallback_correct
        )

        neither_correct = int(
            to_match and not bool(matched_sources & {"Adrien", "ANS", "both"})
        )

        rows.append(
            {
                "Entity ID": entity_id,
                "Stratum ID": stratum_id,
                "reviewed": 1,
                "to_close": int(closes_structure),
                "to_match": int(to_match),
                "initial_correct": initial_correct,
                "initial_na": initial_na,
                "initial_na_correct": initial_na_correct,
                "ans_proposed": ans_proposed,
                "adrien_only_proposed": adrien_only_proposed,
                "adrien_proposed": adrien_proposed,
                "ans_correct": ans_correct,
                "adrien_only_correct": adrien_only_correct,
                "adrien_correct": adrien_correct,
                "ans_fallback_correct": ans_fallback_correct,
                "adrien_only_fallback_correct": adrien_only_fallback_correct,
                "adrien_fallback_correct": adrien_fallback_correct,
                "ans_fallback_mistakes_corrected": ans_fallback_mistakes_corrected,
                "ans_fallback_mistakes_introduced": ans_fallback_mistakes_introduced,
                "adrien_fallback_mistakes_corrected": adrien_fallback_mistakes_corrected,
                "adrien_fallback_mistakes_introduced": adrien_fallback_mistakes_introduced,
                "adrien_only_fallback_mistakes_corrected": adrien_only_fallback_mistakes_corrected,
                "adrien_only_fallback_mistakes_introduced": adrien_only_fallback_mistakes_introduced,
                "neither_correct": neither_correct,
                "with_uncertainty": int(
                    _uncertainty_is_present(selected_uncertainty, config)
                ),
                "business_identified": int(to_match and not selected_is_na),
            }
        )

    return pd.DataFrame(rows)


def analyze_confirmed_decisions(
    confirmed_decisions_df: pd.DataFrame,
    proposal_df: pd.DataFrame,
    *,
    config: ReviewAnalysisConfig = EJ_SIREN_CONFIG,
) -> pd.DataFrame:
    """Return one concise review-analysis row per reviewed stratum.

    Initial correctness is evaluated over entities to match. For ANS and Adrien,
    coverage and correctness-when-proposed are reported separately. Deployable
    correctness applies the production fallback rule entity by entity: use the
    method when it has a proposal (or, for the Adrien-only rule, exactly one
    proposal), otherwise retain Initial. Accepted/satisfactory classifications use
    these deployable fallback point estimates, not Wilson bounds.
    """

    entity_outcomes = build_entity_review_outcomes(
        confirmed_decisions_df, proposal_df, config=config
    )
    names = config.analysis_columns

    count_keys = [
        "reviewed",
        "to_close",
        "to_match",
        "initial_correct",
        "initial_na",
        "initial_na_correct",
        "ans_proposed",
        "adrien_only_proposed",
        "adrien_proposed",
        "ans_correct",
        "adrien_only_correct",
        "adrien_correct",
        "ans_fallback_correct",
        "adrien_only_fallback_correct",
        "adrien_fallback_correct",
        "ans_fallback_mistakes_corrected",
        "ans_fallback_mistakes_introduced",
        "adrien_fallback_mistakes_corrected",
        "adrien_fallback_mistakes_introduced",
        "adrien_only_fallback_mistakes_corrected",
        "adrien_only_fallback_mistakes_introduced",
        "neither_correct",
        "with_uncertainty",
        "business_identified",
    ]

    rows: list[dict[str, object]] = []
    for stratum_id, group in entity_outcomes.groupby("Stratum ID", sort=False):
        counts = {key: int(group[key].sum()) for key in count_keys}
        reviewed = counts["reviewed"]
        to_match = counts["to_match"]

        initial_ci = _wilson_interval_pct(counts["initial_correct"], to_match)
        ans_ci = _wilson_interval_pct(counts["ans_correct"], counts["ans_proposed"])
        adrien_ci = _wilson_interval_pct(
            counts["adrien_correct"], counts["adrien_proposed"]
        )

        initial_pct = _pct(counts["initial_correct"], to_match)
        ans_coverage_pct = _pct(counts["ans_proposed"], to_match)
        adrien_only_coverage_pct = _pct(counts["adrien_only_proposed"], to_match)
        adrien_coverage_pct = _pct(counts["adrien_proposed"], to_match)

        ans_pct = _pct(counts["ans_correct"], counts["ans_proposed"])
        adrien_only_pct = _pct(
            counts["adrien_only_correct"], counts["adrien_only_proposed"]
        )
        adrien_pct = _pct(counts["adrien_correct"], counts["adrien_proposed"])

        ans_fallback_pct = _pct(counts["ans_fallback_correct"], to_match)
        adrien_only_fallback_pct = _pct(
            counts["adrien_only_fallback_correct"], to_match
        )
        adrien_fallback_pct = _pct(counts["adrien_fallback_correct"], to_match)

        def reaches(value: object, threshold: float) -> object:
            if pd.isna(value):
                return pd.NA
            return bool(float(value) >= threshold)

        rows.append(
            {
                names["stratum_id"]: stratum_id,
                names["reviewed"]: reviewed,
                names["to_close"]: counts["to_close"],
                names["to_match"]: to_match,
                names["initial_correct"]: counts["initial_correct"],
                names["initial_na"]: counts["initial_na"],
                names["initial_na_correct"]: counts["initial_na_correct"],
                names["ans_proposed"]: counts["ans_proposed"],
                names["adrien_only_proposed"]: counts["adrien_only_proposed"],
                names["adrien_proposed"]: counts["adrien_proposed"],
                names["ans_correct"]: counts["ans_correct"],
                names["adrien_only_correct"]: counts["adrien_only_correct"],
                names["adrien_correct"]: counts["adrien_correct"],
                names["ans_fallback_correct"]: counts["ans_fallback_correct"],
                names["adrien_only_fallback_correct"]: counts[
                    "adrien_only_fallback_correct"
                ],
                names["adrien_fallback_correct"]: counts["adrien_fallback_correct"],
                names["ans_fallback_mistakes_corrected"]: counts[
                    "ans_fallback_mistakes_corrected"
                ],
                names["ans_fallback_mistakes_introduced"]: counts[
                    "ans_fallback_mistakes_introduced"
                ],
                names["adrien_fallback_mistakes_corrected"]: counts[
                    "adrien_fallback_mistakes_corrected"
                ],
                names["adrien_fallback_mistakes_introduced"]: counts[
                    "adrien_fallback_mistakes_introduced"
                ],
                names["adrien_only_fallback_mistakes_corrected"]: counts[
                    "adrien_only_fallback_mistakes_corrected"
                ],
                names["adrien_only_fallback_mistakes_introduced"]: counts[
                    "adrien_only_fallback_mistakes_introduced"
                ],
                names["neither_correct"]: counts["neither_correct"],
                names["with_uncertainty"]: counts["with_uncertainty"],
                names["business_identified"]: counts["business_identified"],
                names["initial_correct_pct"]: initial_pct,
                names["initial_na_pct"]: _pct(counts["initial_na"], to_match),
                names["initial_na_correct_pct"]: _pct(
                    counts["initial_na_correct"], to_match
                ),
                names["ans_coverage_pct"]: ans_coverage_pct,
                names["adrien_only_coverage_pct"]: adrien_only_coverage_pct,
                names["adrien_coverage_pct"]: adrien_coverage_pct,
                names["ans_correct_pct"]: ans_pct,
                names["adrien_only_correct_pct"]: adrien_only_pct,
                names["adrien_correct_pct"]: adrien_pct,
                names["ans_fallback_correct_pct"]: ans_fallback_pct,
                names["adrien_only_fallback_correct_pct"]: adrien_only_fallback_pct,
                names["adrien_fallback_correct_pct"]: adrien_fallback_pct,
                names["neither_correct_pct"]: _pct(
                    counts["neither_correct"], to_match
                ),
                names["with_uncertainty_pct"]: _pct(
                    counts["with_uncertainty"], reviewed
                ),
                names["to_close_pct"]: _pct(counts["to_close"], reviewed),
                names["business_identified_pct"]: _pct(
                    counts["business_identified"], to_match
                ),
                names["initial_wilson_lower"]: initial_ci[0],
                names["initial_wilson_upper"]: initial_ci[1],
                names["ans_wilson_lower"]: ans_ci[0],
                names["ans_wilson_upper"]: ans_ci[1],
                names["adrien_wilson_lower"]: adrien_ci[0],
                names["adrien_wilson_upper"]: adrien_ci[1],
                names["initial_ge_95"]: reaches(initial_pct, 95.0),
                names["initial_ge_99"]: reaches(initial_pct, 99.0),
                names["ans_ge_95"]: reaches(ans_fallback_pct, 95.0),
                names["ans_ge_99"]: reaches(ans_fallback_pct, 99.0),
                names["adrien_ge_95"]: reaches(adrien_fallback_pct, 95.0),
                names["adrien_ge_99"]: reaches(adrien_fallback_pct, 99.0),
                names["adrien_only_ge_95"]: reaches(
                    adrien_only_fallback_pct, 95.0
                ),
                names["adrien_only_ge_99"]: reaches(
                    adrien_only_fallback_pct, 99.0
                ),
            }
        )

    excluded_design_columns = {
        names["stratum_id"],
        names["population_size"],
        names["sample_size"],
        names["sampling_moe_95"],
        names["sampling_weight"],
        names["first_stage_population_size"],
        names["first_stage_sample_size"],
        names["second_stage_pool_size"],
        names["second_stage_sample_size"],
        names["first_stage_inclusion_probability"],
        names["second_stage_inclusion_probability"],
    }
    analysis_order = [
        names["stratum_id"],
        *[
            name
            for name in config.requested_output_order
            if name not in excluded_design_columns
        ],
    ]
    return pd.DataFrame(rows, columns=analysis_order)


def _numeric_series(df: pd.DataFrame, column: str, *, source: str) -> pd.Series:
    values = pd.to_numeric(df[column], errors="coerce")
    bad = values.isna() & df[column].map(_canonical_text).notna()
    if bad.any():
        examples = df.loc[bad, column].astype(str).head(5).tolist()
        raise DataValidationError(
            f"{source}: column {column!r} must be numeric; examples: {examples}"
        )
    return values.astype(float)


def _require_positive_design(values: pd.Series, *, label: str) -> None:
    if values.isna().any() or (values <= 0).any():
        raise DataValidationError(f"{label} must be positive and non-missing.")


def _add_ege_sampling_design(
    merged: pd.DataFrame,
    *,
    config: ReviewAnalysisConfig,
    names: Mapping[str, str],
    ege_sampling_design_df: pd.DataFrame | None = None,
) -> pd.DataFrame:
    """Compute the two-stage EGE sampling weight from realized design counts."""

    if config.sampling_weight_column and ege_sampling_design_df is not None:
        raise DataValidationError(
            "Provide either sampling_weight_column or ege_sampling_design_df for EGE "
            "analysis, not both."
        )

    if config.sampling_weight_column:
        _raise_if_missing_columns(
            merged, (config.sampling_weight_column,), source="plan_df"
        )
        weight = _numeric_series(
            merged, config.sampling_weight_column, source="plan_df"
        )
        _require_positive_design(weight, label="EGE sampling weight")
        merged[names["sampling_weight"]] = weight.astype("Float64").round(6)
        return merged

    if ege_sampling_design_df is not None:
        required = (
            names["stratum_id"],
            names["first_stage_population_size"],
            names["first_stage_sample_size"],
            names["second_stage_pool_size"],
        )
        _raise_if_missing_columns(
            ege_sampling_design_df, required, source="ege_sampling_design_df"
        )
        design = ege_sampling_design_df[list(required)].copy()
        design["_stratum_id_norm"] = design[names["stratum_id"]].map(
            _canonical_identifier
        )
        if design["_stratum_id_norm"].isna().any():
            raise DataValidationError(
                "ege_sampling_design_df: missing stratum identifiers"
            )
        if design["_stratum_id_norm"].duplicated().any():
            raise DataValidationError(
                "ege_sampling_design_df must contain one row per sampled stratum"
            )
        design = design.drop(columns=[names["stratum_id"]])
        merged = merged.merge(
            design, on="_stratum_id_norm", how="left", validate="1:1"
        )
        ej_population = _numeric_series(
            merged, names["first_stage_population_size"], source="ege_sampling_design_df"
        )
        sampled_ejs = _numeric_series(
            merged, names["first_stage_sample_size"], source="ege_sampling_design_df"
        )
        pooled_eges = _numeric_series(
            merged, names["second_stage_pool_size"], source="ege_sampling_design_df"
        )
    else:
        # Backward-compatible path for workflows that already store the realized
        # two-stage design counts directly in the EGE strata plan.
        required_config = {
            "first_stage_population_size_column": config.first_stage_population_size_column,
            "first_stage_sample_size_column": config.first_stage_sample_size_column,
            "second_stage_pool_size_column": config.second_stage_pool_size_column,
        }
        missing_config = [key for key, value in required_config.items() if not value]
        if missing_config:
            raise DataValidationError(
                "EGE enrichment needs a realized two-stage sampling design. Pass "
                "ege_sampling_design_df (recommended), or configure the legacy "
                "first-stage/second-stage design columns in the EGE strata plan."
            )
        required_columns = [value for value in required_config.values() if value]
        _raise_if_missing_columns(merged, required_columns, source="plan_df")
        ej_population = _numeric_series(
            merged, config.first_stage_population_size_column, source="plan_df"
        )
        sampled_ejs = _numeric_series(
            merged, config.first_stage_sample_size_column, source="plan_df"
        )
        pooled_eges = _numeric_series(
            merged, config.second_stage_pool_size_column, source="plan_df"
        )

    sampled_eges = pd.to_numeric(
        merged[names["reviewed"]], errors="coerce"
    ).fillna(0)

    reviewed_strata = sampled_eges.gt(0)
    for values, label in [
        (ej_population, "EJ population size"),
        (sampled_ejs, "Sampled EJs"),
        (pooled_eges, "Pooled EGEs in sampled EJs"),
    ]:
        if values.loc[reviewed_strata].isna().any() or (
            values.loc[reviewed_strata] <= 0
        ).any():
            raise DataValidationError(
                f"{label} must be positive and non-missing for every stratum with reviewed EGEs."
            )
    _require_positive_design(
        sampled_eges.loc[reviewed_strata], label="Sampled EGEs"
    )

    if (sampled_ejs.loc[reviewed_strata] > ej_population.loc[reviewed_strata]).any():
        raise DataValidationError("Sampled EJs cannot exceed the EJ population size.")
    if (sampled_eges.loc[reviewed_strata] > pooled_eges.loc[reviewed_strata]).any():
        raise DataValidationError(
            "Sampled EGEs cannot exceed the unique pooled EGEs belonging to sampled EJs."
        )

    pi_ej = sampled_ejs / ej_population
    pi_ege = sampled_eges / pooled_eges
    weight = 1.0 / (pi_ej * pi_ege)

    merged[names["first_stage_population_size"]] = ej_population.round().astype("Int64")
    merged[names["first_stage_sample_size"]] = sampled_ejs.round().astype("Int64")
    merged[names["second_stage_pool_size"]] = pooled_eges.round().astype("Int64")
    merged[names["second_stage_sample_size"]] = sampled_eges.round().astype("Int64")
    merged[names["first_stage_inclusion_probability"]] = pi_ej.astype("Float64").round(8)
    merged[names["second_stage_inclusion_probability"]] = pi_ege.astype("Float64").round(8)
    merged[names["sampling_weight"]] = weight.astype("Float64").round(6)
    return merged


def enrich_strata_plan(
    plan_df: pd.DataFrame,
    analysis_df: pd.DataFrame,
    *,
    config: ReviewAnalysisConfig = EJ_SIREN_CONFIG,
    ege_sampling_design_df: pd.DataFrame | None = None,
) -> pd.DataFrame:
    """Join per-stratum analysis onto the plan and add sampling-design fields.

    EJ weights are N_h / n_h unless an explicit weight is supplied. For EGE, the
    recommended path is to pass ``ege_sampling_design_df`` produced by
    :func:`build_ege_two_stage_sampling_design`. EGE weights are then computed as:

        1 / [(sampled EJs / EJ population) *
             (sampled EGEs / unique pooled EGEs in sampled EJs)].

    The completed reviewed EGE count is the second-stage sample size. This keeps
    the realized two-stage design visible in the enriched per-stratum output and
    reuses the same weight for all global reporting.
    """

    names = config.analysis_columns
    _raise_if_missing_columns(
        plan_df,
        (config.stratum_id_column, config.effective_population_size_column),
        source="plan_df",
    )
    _raise_if_missing_columns(analysis_df, (names["stratum_id"],), source="analysis_df")

    # Make re-running enrichment on an already enriched plan safe. Include the
    # hidden standardized sample-size column so older enriched outputs are cleaned.
    standardized_columns = set(config.requested_output_order) | {names["sample_size"]}
    plan = plan_df.drop(
        columns=[c for c in standardized_columns if c in plan_df.columns],
        errors="ignore",
    ).copy()
    plan["_stratum_id_norm"] = plan[config.stratum_id_column].map(_canonical_identifier)
    analysis = analysis_df.copy()
    analysis["_stratum_id_norm"] = analysis[names["stratum_id"]].map(
        _canonical_identifier
    )

    for label, frame in [("analysis_df", analysis), ("plan_df", plan)]:
        duplicate = frame["_stratum_id_norm"].duplicated(keep=False)
        if duplicate.any():
            duplicates = sorted(set(frame.loc[duplicate, "_stratum_id_norm"].dropna()))
            raise DataValidationError(
                f"{label} must contain one row per stratum; duplicates: "
                + ", ".join(duplicates)
            )

    merged = plan.merge(
        analysis.drop(columns=[names["stratum_id"]]),
        on="_stratum_id_norm",
        how="left",
        validate="1:1",
    )

    merged[names["stratum_id"]] = merged[config.stratum_id_column].map(
        _canonical_identifier
    )
    population = _numeric_series(
        merged, config.effective_population_size_column, source="plan_df"
    )
    if population.isna().any() or (population < 0).any():
        raise DataValidationError("Population size must be non-negative and non-missing.")
    merged[names["population_size"]] = population.round().astype("Int64")

    # Sample size is the completed review count unless a plan column is explicitly supplied.
    reviewed_sample = pd.to_numeric(merged[names["reviewed"]], errors="coerce").fillna(0)
    if config.sample_size_column:
        _raise_if_missing_columns(merged, (config.sample_size_column,), source="plan_df")
        sample_size = _numeric_series(merged, config.sample_size_column, source="plan_df")
        mismatch = sample_size.ne(reviewed_sample)
        if mismatch.any():
            examples = merged.loc[
                mismatch,
                [config.stratum_id_column, config.sample_size_column, names["reviewed"]],
            ].head(5)
            raise DataValidationError(
                "The plan sample size must equal the completed review count before global "
                "estimation. Mismatches include:\n" + examples.to_string(index=False)
            )
    else:
        sample_size = reviewed_sample

    # A conservative, descriptive 95% margin of error for a proportion based on
    # examined entities versus the finite stratum population. It is the half-width
    # of a finite-population Wilson-type interval evaluated at p=0.5 (worst-case
    # precision). For EGE's two-stage design it remains a sample-coverage indicator,
    # not a full cluster-design variance estimate; see sampling_precision.py.
    sampling_moe_values: list[object] = []
    for row_index, (n_value, population_value) in enumerate(
        zip(sample_size.tolist(), population.tolist())
    ):
        try:
            moe = sampling_precision.finite_population_margin_of_error_pct(
                n_value, population_value
            )
        except ValueError as exc:
            stratum = merged.iloc[row_index][names["stratum_id"]]
            raise DataValidationError(
                f"Cannot compute sampling margin of error for stratum {stratum!r}: {exc}"
            ) from exc
        sampling_moe_values.append(pd.NA if moe is None else moe)
    merged[names["sampling_moe_95"]] = pd.Series(
        sampling_moe_values, index=merged.index, dtype="Float64"
    )

    # Replace the provisional ordinary Wilson intervals from
    # ``analyze_confirmed_decisions`` with the project's finite-population
    # Wilson-type intervals now that N_h and the realized examined count are known.
    # Initial uses all entities to match. ANS/Adrien conditional correctness uses
    # only entities where that method actually proposed; the FPC still uses the
    # realized examined sample versus the finite stratum population.
    finite_wilson_specs = [
        (
            names["initial_correct"],
            names["to_match"],
            names["initial_wilson_lower"],
            names["initial_wilson_upper"],
            "Initial correctness",
        ),
        (
            names["ans_correct"],
            names["ans_proposed"],
            names["ans_wilson_lower"],
            names["ans_wilson_upper"],
            "ANS correctness when proposed",
        ),
        (
            names["adrien_correct"],
            names["adrien_proposed"],
            names["adrien_wilson_lower"],
            names["adrien_wilson_upper"],
            "Adrien correctness when proposed",
        ),
    ]
    required_wilson_columns = {
        column
        for success_column, trials_column, *_ in finite_wilson_specs
        for column in (success_column, trials_column)
    }
    _raise_if_missing_columns(
        merged,
        sorted(required_wilson_columns),
        source="analysis_df",
    )
    for (
        success_column,
        trials_column,
        lower_column,
        upper_column,
        metric_label,
    ) in finite_wilson_specs:
        success_values = pd.to_numeric(merged[success_column], errors="coerce")
        trial_values = pd.to_numeric(merged[trials_column], errors="coerce")
        lower_values: list[object] = []
        upper_values: list[object] = []
        for row_pos in range(len(merged)):
            stratum = merged.iloc[row_pos][names["stratum_id"]]
            lower, upper = _finite_population_wilson_interval_pct(
                population.iloc[row_pos],
                sample_size.iloc[row_pos],
                trial_values.iloc[row_pos],
                success_values.iloc[row_pos],
                stratum_id=stratum,
                metric_label=metric_label,
            )
            lower_values.append(lower)
            upper_values.append(upper)
        merged[lower_column] = pd.Series(lower_values, index=merged.index, dtype="Float64")
        merged[upper_column] = pd.Series(upper_values, index=merged.index, dtype="Float64")

    # Keep sample_size as an internal series for weighting/validation; do not expose
    # the redundant standardized "Sample size" column in the displayed output.

    if config.is_ege:
        merged = _add_ege_sampling_design(
            merged,
            config=config,
            names=names,
            ege_sampling_design_df=ege_sampling_design_df,
        )
    else:
        if config.sampling_weight_column:
            _raise_if_missing_columns(
                merged, (config.sampling_weight_column,), source="plan_df"
            )
            sampling_weight = _numeric_series(
                merged, config.sampling_weight_column, source="plan_df"
            )
        else:
            sampling_weight = population.div(sample_size.where(sample_size > 0))
        invalid = (population > 0) & (
            sampling_weight.isna() | (sampling_weight <= 0)
        )
        if invalid.any():
            strata = merged.loc[invalid, names["stratum_id"]].astype(str).tolist()
            raise DataValidationError(
                "Every non-empty EJ stratum needs a completed sample; invalid strata: "
                + ", ".join(strata[:10])
            )
        merged[names["sampling_weight"]] = sampling_weight.astype("Float64").round(6)

    count_columns = [
        names["reviewed"],
        names["to_close"],
        names["to_match"],
        names["initial_correct"],
        names["initial_na"],
        names["initial_na_correct"],
        names["ans_proposed"],
        names["adrien_only_proposed"],
        names["adrien_proposed"],
        names["ans_correct"],
        names["adrien_only_correct"],
        names["adrien_correct"],
        names["ans_fallback_correct"],
        names["adrien_only_fallback_correct"],
        names["adrien_fallback_correct"],
        names["ans_fallback_mistakes_corrected"],
        names["ans_fallback_mistakes_introduced"],
        names["adrien_fallback_mistakes_corrected"],
        names["adrien_fallback_mistakes_introduced"],
        names["adrien_only_fallback_mistakes_corrected"],
        names["adrien_only_fallback_mistakes_introduced"],
        names["neither_correct"],
        names["with_uncertainty"],
        names["business_identified"],
    ]
    pct_columns = [
        names["sampling_moe_95"],
        names["initial_correct_pct"],
        names["initial_na_pct"],
        names["initial_na_correct_pct"],
        names["ans_coverage_pct"],
        names["adrien_only_coverage_pct"],
        names["adrien_coverage_pct"],
        names["ans_correct_pct"],
        names["adrien_only_correct_pct"],
        names["adrien_correct_pct"],
        names["ans_fallback_correct_pct"],
        names["adrien_only_fallback_correct_pct"],
        names["adrien_fallback_correct_pct"],
        names["neither_correct_pct"],
        names["with_uncertainty_pct"],
        names["to_close_pct"],
        names["business_identified_pct"],
        names["initial_wilson_lower"],
        names["initial_wilson_upper"],
        names["ans_wilson_lower"],
        names["ans_wilson_upper"],
        names["adrien_wilson_lower"],
        names["adrien_wilson_upper"],
    ]
    bool_columns = [
        names["initial_ge_95"],
        names["initial_ge_99"],
        names["ans_ge_95"],
        names["ans_ge_99"],
        names["adrien_ge_95"],
        names["adrien_ge_99"],
        names["adrien_only_ge_95"],
        names["adrien_only_ge_99"],
    ]

    for column in count_columns:
        merged[column] = pd.to_numeric(merged[column], errors="coerce").fillna(0).astype(int)
    for column in pct_columns:
        merged[column] = pd.to_numeric(merged[column], errors="coerce").astype("Float64")
    for column in bool_columns:
        merged[column] = merged[column].astype("boolean")

    merged = merged.drop(columns=["_stratum_id_norm"], errors="ignore")

    # The enriched output exposes standardized analysis/design columns first. Drop
    # raw plan columns that encode the same design fields; otherwise the default EJ
    # export, for example, ends with a second ``stratum_id`` and ``n_EJ`` after
    # already showing ``Stratum ID`` and ``Population size``.
    redundant_design_sources = {
        source
        for source, standardized in [
            (config.stratum_id_column, names["stratum_id"]),
            (config.effective_population_size_column, names["population_size"]),
            (config.sample_size_column, names["sample_size"]),
            (config.sampling_weight_column, names["sampling_weight"]),
            (
                config.first_stage_population_size_column,
                names["first_stage_population_size"],
            ),
            (
                config.first_stage_sample_size_column,
                names["first_stage_sample_size"],
            ),
            (config.second_stage_pool_size_column, names["second_stage_pool_size"]),
            (
                config.second_stage_sample_size_column,
                names["second_stage_sample_size"],
            ),
        ]
        if source and source != standardized and source in merged.columns
    }
    merged = merged.drop(columns=sorted(redundant_design_sources), errors="ignore")

    front = [column for column in config.requested_output_order if column in merged.columns]
    remaining = [column for column in merged.columns if column not in front]
    return merged[front + remaining]


def _excel_display_header(column: object, all_headers: set[str]) -> str:
    """Return a compact display header for Excel without unit-only ``(%)`` suffixes.

    Percentage columns in the analysis can have a count column with the same base
    name (for example ``Initial correct`` and ``Initial correct (%)``). In that
    case, use ``rate`` to keep the exported headers unambiguous.
    """

    header = str(column)
    suffix = " (%)"
    if not header.endswith(suffix):
        return header

    base = header[: -len(suffix)]
    return f"{base} rate" if base in all_headers else base


def _excel_threshold_fraction(value: object) -> object:
    """Normalize a threshold such as ``95%`` or ``95`` to the Excel fraction 0.95."""

    if pd.isna(value):
        return pd.NA
    if isinstance(value, str):
        text = value.strip()
        if text.endswith("%"):
            text = text[:-1].strip()
        try:
            number = float(text)
        except ValueError:
            return value
    else:
        try:
            number = float(value)
        except (TypeError, ValueError):
            return value
    return number / 100.0 if abs(number) > 1 else number


def _drop_redundant_excel_columns(df: pd.DataFrame) -> pd.DataFrame:
    """Remove raw design columns duplicated by standardized analysis columns.

    ``enrich_strata_plan`` already removes configured source columns, but this
    export-time guard also cleans older/pre-enriched dataframes that still carry
    trailing ``stratum_id`` and ``n_EJ`` / ``n_EGE`` columns.
    """

    exported = df.copy()
    headers = {str(column).strip(): column for column in exported.columns}
    to_drop: list[object] = []

    if "Stratum ID" in headers:
        to_drop.extend(
            column
            for column in exported.columns
            if str(column).strip().lower() == "stratum_id"
        )

    if "Population size" in headers:
        to_drop.extend(
            column
            for column in exported.columns
            if re.fullmatch(r"n_(?:ej|ege)", str(column).strip(), flags=re.IGNORECASE)
        )

    return exported.drop(columns=list(dict.fromkeys(to_drop)), errors="ignore")


def _wrap_excel_header(
    header: object,
    *,
    max_lines: int = 4,
    minimum_width: int = 8,
) -> tuple[str, int, int]:
    """Wrap a header at word boundaries using no more than ``max_lines``.

    The returned width is the narrowest practical Excel column width that keeps
    the header within the requested line count. Explicit line breaks are written
    into the worksheet so rendering does not depend on Excel's auto-wrap heuristics.
    """

    text = " ".join(str(header).split())
    if not text:
        return "", minimum_width, 1

    words = text.split()
    longest_word = max((len(word) for word in words), default=minimum_width)
    width = max(minimum_width, longest_word)

    while True:
        lines = textwrap.wrap(
            text,
            width=width,
            break_long_words=False,
            break_on_hyphens=False,
        ) or [text]
        if len(lines) <= max_lines:
            return "\n".join(lines), width, len(lines)
        width += 1


def _prepare_dataframe_for_excel(
    df: pd.DataFrame,
) -> tuple[pd.DataFrame, dict[str, str]]:
    """Prepare an export-only dataframe and per-column Excel number formats.

    Analysis dataframes keep percentages in percentage points (for example 82.35)
    because threshold logic is expressed as 95 / 99. The Excel copy converts those
    values to true Excel fractions (0.8235) and relies on cell number formats to
    display ``82.35%``. This keeps spreadsheet values numeric and calculations safe.
    """

    exported = _drop_redundant_excel_columns(df)
    all_headers = {str(column) for column in exported.columns}
    rename_map: dict[object, str] = {}
    number_formats_by_original: dict[object, str] = {}

    for column in exported.columns:
        header = str(column)

        if header.endswith(" (%)"):
            exported[column] = pd.to_numeric(exported[column], errors="coerce") / 100.0
            rename_map[column] = _excel_display_header(column, all_headers)
            number_formats_by_original[column] = "0.00%"
            continue

        # Wilson bounds and the sampling margin of error are also percentages, but
        # the 95% in their headers denotes the confidence level rather than a unit
        # suffix, so keep that text while converting values to Excel fractions.
        if "Wilson 95%" in header or header == "Sampling margin of error — 95%":
            exported[column] = pd.to_numeric(exported[column], errors="coerce") / 100.0
            number_formats_by_original[column] = "0.00%"
            continue

        # Inclusion probabilities are already stored as fractions.
        if header.endswith("inclusion probability"):
            number_formats_by_original[column] = "0.00%"
            continue

        if header == "Threshold":
            exported[column] = exported[column].map(_excel_threshold_fraction)
            number_formats_by_original[column] = "0%"
            continue

        if header == "Sampling weight":
            number_formats_by_original[column] = "#,##0.00"
            continue

        if "percentage points" in header.lower():
            number_formats_by_original[column] = "0.00"
            continue

        if pd.api.types.is_bool_dtype(exported[column].dtype):
            continue

        if pd.api.types.is_numeric_dtype(exported[column].dtype):
            numeric = pd.to_numeric(exported[column], errors="coerce")
            finite = numeric.dropna()
            if finite.empty or finite.mod(1).eq(0).all():
                number_formats_by_original[column] = "#,##0"
            else:
                # Show at most two decimals and suppress unnecessary trailing zeros.
                number_formats_by_original[column] = "#,##0.##"

    exported = exported.rename(columns=rename_map)
    number_formats = {
        rename_map.get(column, str(column)): fmt
        for column, fmt in number_formats_by_original.items()
    }
    return exported, number_formats


def _excel_display_length_for_series(
    series: pd.Series, number_format: str | None
) -> int:
    """Estimate the rendered width of a numeric/boolean Excel column.

    This deliberately uses the cell's display format rather than ``str(float)`` so
    binary floating-point tails do not make compact percentage columns artificially
    wide.
    """

    if pd.api.types.is_bool_dtype(series.dtype):
        return 5  # FALSE

    numeric = pd.to_numeric(series, errors="coerce").dropna().head(250)
    if numeric.empty:
        return 0

    def render(value: float) -> str:
        if number_format == "0.00%":
            return f"{value * 100:.2f}%"
        if number_format == "0%":
            return f"{value * 100:.0f}%"
        if number_format == "#,##0.00":
            return f"{value:,.2f}"
        if number_format == "#,##0":
            return f"{value:,.0f}"
        if number_format == "0.00":
            return f"{value:.2f}"
        if number_format == "#,##0.##":
            return f"{value:,.2f}".rstrip("0").rstrip(".")
        return f"{value:g}"

    return max(len(render(float(value))) for value in numeric)


def _format_excel_worksheet(
    worksheet,
    exported_df: pd.DataFrame,
    number_formats: Mapping[str, str],
    *,
    header_row: int = 1,
    table_end_row: int | None = None,
    compact_numeric_columns: bool = False,
) -> None:
    """Apply restrained, consistent formatting to an exported dataframe sheet.

    ``compact_numeric_columns`` is intended for wide per-stratum analysis tables.
    Long headers are explicitly wrapped at word boundaries to a maximum of four
    lines, while the column width remains driven mainly by the compact cell values.
    """

    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    if table_end_row is None:
        table_end_row = header_row + len(exported_df)

    first_data_row = header_row + 1
    last_column = len(exported_df.columns)

    for cell in worksheet[header_row]:
        if cell.column > last_column:
            break
        cell.font = Font(bold=True)
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )

    if header_row == 1:
        worksheet.freeze_panes = "A2"

    if last_column and table_end_row >= header_row:
        last_letter = get_column_letter(last_column)
        worksheet.auto_filter.ref = f"A{header_row}:{last_letter}{table_end_row}"

    max_header_lines = 1
    for column_index, column in enumerate(exported_df.columns, start=1):
        letter = get_column_letter(column_index)
        series = exported_df[column]
        header = str(column)

        is_numeric = (
            pd.api.types.is_numeric_dtype(series.dtype)
            and not pd.api.types.is_bool_dtype(series.dtype)
        )
        is_boolean = pd.api.types.is_bool_dtype(series.dtype)

        number_format = number_formats.get(header)
        if is_numeric or is_boolean:
            max_content = _excel_display_length_for_series(series, number_format)
        else:
            content_lengths = [len(str(value)) for value in series.dropna().head(250)]
            max_content = max(content_lengths, default=0)

        if compact_numeric_columns:
            wrapped_header, header_width, header_lines = _wrap_excel_header(
                header, max_lines=4, minimum_width=8
            )
            worksheet.cell(row=header_row, column=column_index).value = wrapped_header
            max_header_lines = max(max_header_lines, header_lines)

            if is_numeric or is_boolean:
                # Keep analytical numeric columns narrow: the smallest width that
                # holds the values and keeps the header at four lines or fewer.
                width = max(8, max_content + 2, header_width + 1)
                width = min(width, 22)
            else:
                # Identifier/text columns can be slightly wider, but still benefit
                # from the explicit four-line header wrapping.
                width = max(10, max_content + 2, header_width + 1)
                width = min(width, 30)
        elif is_numeric:
            width_cap = 18
            header_hint = min(len(header), 24)
            width = min(max(10, max(max_content, header_hint) + 2), width_cap)
        else:
            width_cap = 40
            header_hint = min(len(header), 24)
            width = min(max(10, max(max_content, header_hint) + 2), width_cap)

        worksheet.column_dimensions[letter].width = width

        if number_format and table_end_row >= first_data_row:
            for row in range(first_data_row, table_end_row + 1):
                worksheet.cell(row=row, column=column_index).number_format = number_format

    if compact_numeric_columns:
        # About 15 points per wrapped line plus a small top/bottom margin.
        worksheet.row_dimensions[header_row].height = min(66, 6 + 15 * max_header_lines)
    else:
        worksheet.row_dimensions[header_row].height = 42


def export_dataframe_to_excel(
    df: pd.DataFrame,
    output_path: str | Path,
    *,
    sheet_name: str = "Sheet1",
    supplementary_sheets: Mapping[str, pd.DataFrame] | None = None,
) -> None:
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    exported_df, number_formats = _prepare_dataframe_for_excel(df)
    supplementary_sheets = supplementary_sheets or {}
    if sheet_name in supplementary_sheets:
        raise DataValidationError(
            f"Supplementary sheet name {sheet_name!r} duplicates the main sheet."
        )
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        exported_df.to_excel(
            writer, index=False, sheet_name=sheet_name, na_rep=""
        )
        _format_excel_worksheet(
            writer.sheets[sheet_name],
            exported_df,
            number_formats,
            table_end_row=1 + len(exported_df),
            compact_numeric_columns=(sheet_name == "strata_plan_review_analysis"),
        )
        for extra_sheet_name, extra_frame in supplementary_sheets.items():
            extra_exported, extra_formats = _prepare_dataframe_for_excel(extra_frame)
            extra_exported.to_excel(
                writer,
                index=False,
                sheet_name=extra_sheet_name,
                na_rep="",
            )
            _format_excel_worksheet(
                writer.sheets[extra_sheet_name],
                extra_exported,
                extra_formats,
                table_end_row=1 + len(extra_exported),
            )


def export_dataframe_to_parquet(df: pd.DataFrame, output_path: str | Path) -> None:
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    df.to_parquet(output, index=False)


def export_dataframe(
    df: pd.DataFrame,
    output_path: str | Path,
    *,
    sheet_name: str = "Sheet1",
    supplementary_sheets: Mapping[str, pd.DataFrame] | None = None,
) -> None:
    output = Path(output_path)
    suffix = output.suffix.lower()
    if suffix in _VALID_EXCEL_SUFFIXES:
        export_dataframe_to_excel(
            df,
            output,
            sheet_name=sheet_name,
            supplementary_sheets=supplementary_sheets,
        )
    elif suffix == ".parquet":
        if supplementary_sheets:
            raise DataValidationError(
                "Supplementary sheets are supported only for Excel exports."
            )
        export_dataframe_to_parquet(df, output)
    else:
        raise DataValidationError(
            f"Unsupported export format for {output.name!r}; expected Excel or parquet"
        )


def run_analysis(
    *,
    confirmed_decisions_path: str | Path,
    proposal_path: str | Path,
    plan_path: str | Path,
    config: ReviewAnalysisConfig = EJ_SIREN_CONFIG,
) -> dict[str, pd.DataFrame]:
    """Run per-stratum review analysis and enrich the strata plan."""

    confirmed_decisions_df = load_confirmed_decisions(
        confirmed_decisions_path, config=config
    )
    proposal_df = load_proposal_table(proposal_path, config=config)
    plan_df = load_strata_plan(plan_path, config=config)

    analysis_df = analyze_confirmed_decisions(
        confirmed_decisions_df, proposal_df, config=config
    )
    enriched_plan_df = enrich_strata_plan(plan_df, analysis_df, config=config)

    return {
        "confirmed_decisions_df": confirmed_decisions_df,
        "proposal_df": proposal_df,
        "plan_df": plan_df,
        "analysis_df": analysis_df,
        "enriched_plan_df": enriched_plan_df,
    }


def main(argv: Sequence[str] | None = None) -> int:
    """Command-line interface for per-stratum review analysis."""

    import argparse

    parser = argparse.ArgumentParser(
        description=(
            "Analyze confirmed review decisions by stratum and enrich a strata plan "
            "with descriptive statistics, Wilson intervals, a finite-population Wilson sampling margin of error and sampling-design fields."
        )
    )
    parser.add_argument("--confirmed-decisions-path", required=True)
    parser.add_argument("--proposal-path", required=True)
    parser.add_argument("--plan-path", required=True)
    parser.add_argument("--enriched-plan-out", required=True)

    parser.add_argument("--entity-name", default="EJ")
    parser.add_argument("--business-name", default="siren")
    parser.add_argument("--entity-id-column", default="EJ")
    parser.add_argument("--business-id-proposal-column", default="siren_proposal")
    parser.add_argument("--reviewed-business-id-column", default="Siren_retenu")
    parser.add_argument("--proposal-source-column", default="proposal_source")
    parser.add_argument("--stratum-id-column", default="stratum_id")
    parser.add_argument("--uncertainty-column", default="Incertitude1")
    parser.add_argument("--adrien-status-column", default="statut_validation_Adrien")
    parser.add_argument("--initial-proposal-flag-column")
    parser.add_argument("--population-size-column")
    parser.add_argument("--sample-size-column")
    parser.add_argument("--sampling-weight-column")
    parser.add_argument("--first-stage-population-size-column")
    parser.add_argument("--first-stage-sample-size-column")
    parser.add_argument("--second-stage-pool-size-column")
    parser.add_argument("--second-stage-sample-size-column")
    args = parser.parse_args(argv)

    config = ReviewAnalysisConfig(
        entity_name=args.entity_name,
        business_name=args.business_name,
        entity_id_column=args.entity_id_column,
        business_id_proposal_column=args.business_id_proposal_column,
        reviewed_business_id_column=args.reviewed_business_id_column,
        proposal_source_column=args.proposal_source_column,
        stratum_id_column=args.stratum_id_column,
        uncertainty_column=args.uncertainty_column,
        adrien_status_column=args.adrien_status_column,
        initial_proposal_flag_column=args.initial_proposal_flag_column,
        population_size_column=args.population_size_column,
        sample_size_column=args.sample_size_column,
        sampling_weight_column=args.sampling_weight_column,
        first_stage_population_size_column=args.first_stage_population_size_column,
        first_stage_sample_size_column=args.first_stage_sample_size_column,
        second_stage_pool_size_column=args.second_stage_pool_size_column,
        second_stage_sample_size_column=args.second_stage_sample_size_column,
    )

    outputs = run_analysis(
        confirmed_decisions_path=args.confirmed_decisions_path,
        proposal_path=args.proposal_path,
        plan_path=args.plan_path,
        config=config,
    )
    export_dataframe(
        outputs["enriched_plan_df"],
        args.enriched_plan_out,
        sheet_name="strata_plan_review_analysis",
    )

    print("Per-stratum analysis completed successfully.")
    return 0


if __name__ == "__main__":  # pragma: no cover
    raise SystemExit(main())
