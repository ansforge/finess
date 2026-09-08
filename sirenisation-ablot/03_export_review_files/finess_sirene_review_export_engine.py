"""Shared FINESS/SIRENE review-export engine.

The core review logic works with the neutral internal identifiers ``entity_id``
and ``business_id``.  A :class:`ReviewPipelineConfig` maps the real dataframe
columns (for example ``EJ``/``siren_proposal`` or ``EGE``/``siret_proposal``)
to those internal names.  Dataframes returned to callers always retain their
original/public column names.
"""

from __future__ import annotations

import re
import shutil
import warnings
from collections.abc import Callable, Mapping, Sequence
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

_INTERNAL_ENTITY_ID = "entity_id"
_INTERNAL_BUSINESS_ID = "business_id"


@dataclass(frozen=True)
class HyperlinkSpec:
    """Describe an optional hyperlink rendered in one output column.

    ``label`` and ``url`` receive a neutral row context containing
    ``entity_id`` and ``business_id`` plus every public source column.
    ``when`` can suppress the link for selected rows.
    """

    label: Callable[[Mapping[str, Any]], str]
    url: Callable[[Mapping[str, Any]], str]
    when: Callable[[Mapping[str, Any]], bool] | None = None


@dataclass(frozen=True)
class ReviewPipelineConfig:
    """Configuration for one FINESS/SIRENE review pipeline."""

    # Human-readable labels used in messages/docstrings.
    entity_label: str
    business_label: str

    # Real/public dataframe column names.  Core logic maps these temporarily
    # to ``entity_id`` and ``business_id``.
    entity_id_col: str
    business_id_col: str

    identification_cols: Sequence[str]
    review_cols: Sequence[str]
    info_cols: Sequence[str]
    strat_cols: Sequence[str]
    context_cols: Sequence[str] = ()

    # ``assignment_col`` is metadata from the strata plan, not a required
    # proposals column.  The engine attaches/overwrites it from the plan before
    # splitting outputs, making the plan the single source of truth.
    assignment_col: str = "reviewer"
    stratum_col: str = "stratum_id"
    sample_stratum_col: str = "stratum_id"
    sample_size_col: str = "sample_size"
    random_key_col: str = "random_key"
    candidate_rank_col: str = "candidate_rank"
    required_cols: Sequence[str] = ()

    date_cols: frozenset[str] = field(default_factory=frozenset)
    centered_cols: frozenset[str] = field(default_factory=frozenset)

    # Prefixes drive the subtle EGE-style column shading and width rules.
    entity_prefixes: tuple[str, ...] = ()
    business_prefixes: tuple[str, ...] = ()
    context_prefixes: tuple[str, ...] = ()

    width_overrides: Mapping[str, float] = field(default_factory=dict)
    hyperlinks: Mapping[str, HyperlinkSpec] = field(default_factory=dict)

    display_total: int = 20

    # Default workbook formatting.
    base_entity_colors: tuple[str, str] = ("FDE9D9", "FFF2CC")
    review_fill: str = "EAF3FF"
    header_fill: str = "2F3A4A"
    header_font_color: str = "FFFFFF"
    thin_gray: str = "B7B7B7"
    medium_gray: str = "8A8A8A"
    date_number_format: str = "yyyy-mm-dd"

    @property
    def column_order(self) -> list[str]:
        return (
            list(self.identification_cols)
            + list(self.review_cols)
            + list(self.info_cols)
            + list(self.context_cols)
            + list(self.strat_cols)
        )


# ---------------------------------------------------------------------------
# Neutral/internal dataframe handling
# ---------------------------------------------------------------------------
def _validate_internal_name_collisions(df: pd.DataFrame, config: ReviewPipelineConfig) -> None:
    rename_sources = {config.entity_id_col, config.business_id_col}
    for internal_name in (_INTERNAL_ENTITY_ID, _INTERNAL_BUSINESS_ID):
        if internal_name in df.columns and internal_name not in rename_sources:
            raise ValueError(
                f"Input already contains reserved internal column {internal_name!r}. "
                "Rename it or configure that column as the corresponding identifier."
            )


def _to_internal(df: pd.DataFrame, config: ReviewPipelineConfig) -> pd.DataFrame:
    """Return a copy whose configured IDs use neutral core names."""
    _validate_internal_name_collisions(df, config)
    rename_map: dict[str, str] = {}
    if config.entity_id_col != _INTERNAL_ENTITY_ID:
        rename_map[config.entity_id_col] = _INTERNAL_ENTITY_ID
    if config.business_id_col != _INTERNAL_BUSINESS_ID:
        rename_map[config.business_id_col] = _INTERNAL_BUSINESS_ID
    return df.rename(columns=rename_map).copy()


def _to_public(df: pd.DataFrame, config: ReviewPipelineConfig) -> pd.DataFrame:
    """Map neutral core identifiers back to configured/public column names."""
    rename_map: dict[str, str] = {}
    if config.entity_id_col != _INTERNAL_ENTITY_ID:
        rename_map[_INTERNAL_ENTITY_ID] = config.entity_id_col
    if config.business_id_col != _INTERNAL_BUSINESS_ID:
        rename_map[_INTERNAL_BUSINESS_ID] = config.business_id_col
    return df.rename(columns=rename_map).copy()


# ---------------------------------------------------------------------------
# Generic helpers
# ---------------------------------------------------------------------------
def safe_name(value: Any) -> str:
    """Make a value safe for use as a folder/file name."""
    if pd.isna(value):
        return "unknown"

    s = str(value).strip()
    s = re.sub(r"[^\w.\-]+", "_", s, flags=re.UNICODE)
    s = re.sub(r"_+", "_", s).strip("_")
    return s or "unknown"


def normalize_identifier(value: Any, width: int | None = None) -> str:
    """Normalize an identifier for matching without altering displayed values."""
    if pd.isna(value):
        return ""

    s = str(value).strip()
    if re.fullmatch(r"\d+\.0", s):
        s = s[:-2]
    if width is not None and s.isdigit():
        s = s.zfill(width)
    return s


def _normalized_strata_plan(
    strata_plan: pd.DataFrame,
    config: ReviewPipelineConfig,
) -> tuple[pd.DataFrame, dict[str, int], dict[str, Any]]:
    """Normalize strata-plan metadata used for sampling and assignment.

    Assignment is authoritative in the strata plan.  Proposal-side assignment
    values are deliberately ignored by :func:`export_review_files`.
    """
    required = {
        config.sample_stratum_col,
        config.sample_size_col,
        config.assignment_col,
    }
    missing = required - set(strata_plan.columns)
    if missing:
        raise ValueError(
            "strata_plan must contain columns: "
            f"{sorted(required)}. Missing: {sorted(missing)}"
        )

    plan = strata_plan[
        [config.sample_stratum_col, config.sample_size_col, config.assignment_col]
    ].copy()
    plan["_stratum_match_key"] = plan[config.sample_stratum_col].map(
        normalize_identifier
    )
    sample_size = pd.to_numeric(
        plan[config.sample_size_col],
        errors="coerce",
    )
    invalid_sample_size = (
        sample_size.isna()
        | sample_size.lt(0)
        | sample_size.mod(1).ne(0)
    )
    if invalid_sample_size.any():
        bad_values = plan.loc[
            invalid_sample_size,
            [config.sample_stratum_col, config.sample_size_col],
        ].head(10).to_dict("records")
        raise ValueError(
            f"{config.sample_size_col!r} must contain non-negative integers. "
            f"Invalid values: {bad_values}"
        )
    plan[config.sample_size_col] = sample_size.astype("int64")

    missing_assignment = plan[config.assignment_col].map(
        lambda value: pd.isna(value) or str(value).strip() == ""
    )
    if missing_assignment.any():
        bad_strata = plan.loc[
            missing_assignment, config.sample_stratum_col
        ].astype(str).tolist()
        raise ValueError(
            f"Strata plan contains missing {config.assignment_col!r} values for "
            f"strata: {bad_strata}"
        )

    duplicate_strata = plan.loc[
        plan["_stratum_match_key"].duplicated(keep=False),
        "_stratum_match_key",
    ].drop_duplicates()
    if not duplicate_strata.empty:
        raise ValueError(
            "The final strata plan must contain exactly one row per stratum. "
            f"Duplicate normalized strata: {duplicate_strata.head(20).tolist()}"
        )

    unique_plan = plan
    sample_map = unique_plan.set_index("_stratum_match_key")[
        config.sample_size_col
    ].to_dict()
    assignment_map = unique_plan.set_index("_stratum_match_key")[
        config.assignment_col
    ].to_dict()
    return plan, sample_map, assignment_map


def _attach_assignment_from_strata_plan(
    df: pd.DataFrame,
    strata_plan: pd.DataFrame,
    config: ReviewPipelineConfig,
) -> pd.DataFrame:
    """Attach authoritative assignment metadata by normalized stratum ID.

    If the proposals dataframe happens to contain an assignment column, it is
    discarded first and never consulted.
    """
    _, _, assignment_map = _normalized_strata_plan(strata_plan, config)

    out = df.drop(columns=[config.assignment_col], errors="ignore").copy()
    match_keys = out[config.stratum_col].map(normalize_identifier)
    missing_keys = sorted(
        {key for key in pd.unique(match_keys) if key not in assignment_map}
    )
    if missing_keys:
        raise ValueError(
            "Proposal strata are missing from the strata plan, so assignment "
            f"cannot be determined: {missing_keys}"
        )

    out[config.assignment_col] = match_keys.map(assignment_map)
    return out


def _validate_source_df(df: pd.DataFrame, config: ReviewPipelineConfig) -> None:
    required = {
        config.entity_id_col,
        config.business_id_col,
        config.stratum_col,
        config.random_key_col,
        config.candidate_rank_col,
        *config.required_cols,
    }
    missing = required - set(df.columns)
    if missing:
        raise ValueError(f"Missing required columns: {sorted(missing)}")

    if config.display_total < 0:
        raise ValueError("display_total must be greater than or equal to zero.")


def add_empty_review_columns(
    df: pd.DataFrame,
    config: ReviewPipelineConfig,
) -> pd.DataFrame:
    """Ensure manual-review columns exist without changing existing values."""
    out = df.copy()
    for col in config.review_cols:
        if col not in out.columns:
            out[col] = pd.NA
    return out


def ordered_columns(df: pd.DataFrame, config: ReviewPipelineConfig) -> list[str]:
    """Preferred configured order, followed by any unexpected source columns."""
    cols = [c for c in config.column_order if c in df.columns]
    extras = [c for c in df.columns if c not in cols]
    return cols + extras


def apply_confirmed_decisions(
    df: pd.DataFrame,
    confirmed_decisions: pd.DataFrame | None,
    config: ReviewPipelineConfig,
) -> pd.DataFrame:
    """Prefill confirmed review decisions on the first row of each entity.

    ``confirmed_decisions`` is optional and must contain exactly one logical row
    per entity (additional columns are allowed).  The required columns are the
    configured entity identifier plus every configured review column; for
    example ``EJ, Siren_retenu, Incertitude1, Commentaire`` or
    ``EGE, Siret_retenu, Incertitude1, Commentaire``.

    Matching uses :func:`normalize_identifier`, while displayed identifiers are
    left untouched.  For a matched entity, review values are cleared from all
    of its candidate rows and the confirmed values are written only on the
    first row in canonical review order.  This makes the confirmed-decision
    database authoritative for matched entities without changing sampling.
    """
    if confirmed_decisions is None:
        return df.copy()

    required = {config.entity_id_col, *config.review_cols}
    missing = required - set(confirmed_decisions.columns)
    if missing:
        raise ValueError(
            "confirmed_decisions must contain columns: "
            f"{sorted(required)}. Missing: {sorted(missing)}"
        )

    decisions = confirmed_decisions[
        [config.entity_id_col, *config.review_cols]
    ].copy()
    decisions["_entity_match_key"] = decisions[config.entity_id_col].map(
        normalize_identifier
    )

    missing_id = decisions["_entity_match_key"].eq("")
    if missing_id.any():
        # Report spreadsheet-like data-row positions independently of the
        # dataframe's index type.
        bad_rows = [pos + 2 for pos, is_missing in enumerate(missing_id) if is_missing]
        raise ValueError(
            "confirmed_decisions contains missing/blank entity identifiers at "
            f"data rows: {bad_rows}"
        )

    duplicate_key = decisions["_entity_match_key"].duplicated(keep=False)
    if duplicate_key.any():
        duplicate_ids = decisions.loc[
            duplicate_key, config.entity_id_col
        ].astype(str).tolist()
        raise ValueError(
            "confirmed_decisions must contain one row per entity after identifier "
            f"normalization. Duplicates: {duplicate_ids}"
        )

    out = add_empty_review_columns(df, config)
    out = order_rows_for_review(out, config)
    entity_keys = out[config.entity_id_col].map(normalize_identifier)

    decision_by_key = decisions.set_index("_entity_match_key")
    matched_keys = set(entity_keys).intersection(decision_by_key.index)
    if not matched_keys:
        return out

    matched_rows = entity_keys.isin(matched_keys)
    for col in config.review_cols:
        # Review columns are manual-entry fields and may mix strings, numbers,
        # booleans and blanks, so use object dtype before authoritative writes.
        out[col] = out[col].astype("object")
        out.loc[matched_rows, col] = pd.NA

    first_row_mask = matched_rows & ~entity_keys.duplicated(keep="first")
    for idx in out.index[first_row_mask]:
        key = entity_keys.loc[idx]
        for col in config.review_cols:
            out.at[idx, col] = decision_by_key.at[key, col]

    return out


def order_rows_for_review(
    df: pd.DataFrame,
    config: ReviewPipelineConfig,
) -> pd.DataFrame:
    """Apply the deterministic review ordering using neutral internal identifiers.

    Rows are sorted stably by random key, entity ID, then candidate rank.  The
    candidate rank is sorted numerically when possible; this fixes the dormant
    EGE helper bug where a numeric sort key was computed but never used.
    """
    internal = _to_internal(df, config)

    random_key_internal = config.random_key_col
    rank_internal = config.candidate_rank_col
    internal["_candidate_rank_order"] = pd.to_numeric(
        internal[rank_internal], errors="coerce"
    )

    # Preserve EGE's primary ordering.  Numeric rank matters only within an
    # entity; mergesort keeps ties/source order stable.
    internal = internal.sort_values(
        by=[random_key_internal, _INTERNAL_ENTITY_ID, "_candidate_rank_order"],
        kind="mergesort",
        na_position="last",
    ).reset_index(drop=True)
    internal.drop(columns=["_candidate_rank_order"], inplace=True)

    return _to_public(internal, config)


def sample_first_unique_entities_by_stratum(
    df: pd.DataFrame,
    strata_plan: pd.DataFrame,
    config: ReviewPipelineConfig,
    display_total: int | None = None,
) -> tuple[pd.DataFrame, int]:
    """Select review/context groups for exactly one stratum.

    The selection is entity-level for both supported cases: EJ groups for the
    legal-entity review and EGE groups for the establishment review.  Every
    candidate business row belonging to a retained entity is kept.
    """
    required_df_cols = {config.stratum_col, config.entity_id_col, config.business_id_col}
    missing_df_cols = required_df_cols - set(df.columns)
    if missing_df_cols:
        raise ValueError(
            f"df must contain columns: {sorted(required_df_cols)}. "
            f"Missing: {sorted(missing_df_cols)}"
        )

    _, sample_map, _ = _normalized_strata_plan(strata_plan, config)

    display_total = config.display_total if display_total is None else display_total
    if display_total < 0:
        raise ValueError("display_total must be greater than or equal to zero.")

    stratum_values = pd.unique(df[config.stratum_col])
    if len(stratum_values) != 1:
        raise ValueError(
            "sample_first_unique_entities_by_stratum must receive exactly one stratum."
        )

    stratum_id = stratum_values[0]
    requested_review_n = int(sample_map.get(normalize_identifier(stratum_id), 0))

    ordered_public = order_rows_for_review(df, config)
    ordered = _to_internal(ordered_public, config)
    entity_order = pd.unique(ordered[_INTERNAL_ENTITY_ID])
    available_entity_n = len(entity_order)

    if requested_review_n > available_entity_n:
        warnings.warn(
            (
                f"Stratum {stratum_id!r}: requested sample size is "
                f"{requested_review_n}, but only {available_entity_n} unique "
                f"{config.entity_label} values are available. All available "
                f"{config.entity_label} values will be included."
            ),
            category=UserWarning,
            stacklevel=2,
        )

    review_entity_count = min(requested_review_n, available_entity_n)
    display_entity_count = min(
        max(display_total, review_entity_count),
        available_entity_n,
    )

    selected_entities = entity_order[:display_entity_count]
    sampled = ordered[ordered[_INTERNAL_ENTITY_ID].isin(selected_entities)].copy()
    return _to_public(sampled, config), review_entity_count


# ---------------------------------------------------------------------------
# Excel workbook formatting
# ---------------------------------------------------------------------------
def _to_argb(hex_color: str) -> str:
    return "FF" + hex_color.replace("#", "").upper()


def _blend_with_white(hex_color: str, amount: float) -> str:
    hex_color = hex_color.replace("#", "").upper()
    r = int(hex_color[0:2], 16)
    g = int(hex_color[2:4], 16)
    b = int(hex_color[4:6], 16)
    r = round(r + (255 - r) * amount)
    g = round(g + (255 - g) * amount)
    b = round(b + (255 - b) * amount)
    return f"{r:02X}{g:02X}{b:02X}"


def _blend_with_black(hex_color: str, amount: float) -> str:
    hex_color = hex_color.replace("#", "").upper()
    r = int(hex_color[0:2], 16)
    g = int(hex_color[2:4], 16)
    b = int(hex_color[4:6], 16)
    r = round(r * (1 - amount))
    g = round(g * (1 - amount))
    b = round(b * (1 - amount))
    return f"{r:02X}{g:02X}{b:02X}"


def _cell_to_text(value: Any) -> str:
    if pd.isna(value):
        return ""
    if isinstance(value, pd.Timestamp):
        return value.strftime("%Y-%m-%d")
    return str(value)


def _starts_with_any(value: str, prefixes: tuple[str, ...]) -> bool:
    return bool(prefixes) and value.startswith(prefixes)


def _column_fill(col_name: str, base_color: str, config: ReviewPipelineConfig) -> str:
    if col_name in config.review_cols:
        return config.review_fill
    if _starts_with_any(col_name, config.entity_prefixes):
        return _blend_with_black(base_color, 0.04)
    if _starts_with_any(col_name, config.business_prefixes):
        return _blend_with_white(base_color, 0.04)
    if _starts_with_any(col_name, config.context_prefixes):
        return _blend_with_white(base_color, 0.10)
    if col_name in config.strat_cols:
        return _blend_with_white(base_color, 0.04)
    return base_color


def _neutral_row_context(row: pd.Series, config: ReviewPipelineConfig) -> dict[str, Any]:
    context = row.to_dict()
    context[_INTERNAL_ENTITY_ID] = row.get(config.entity_id_col)
    context[_INTERNAL_BUSINESS_ID] = row.get(config.business_id_col)
    return context


def _hyperlink_value(
    col_name: str,
    row: pd.Series,
    config: ReviewPipelineConfig,
) -> tuple[str | None, str | None]:
    spec = config.hyperlinks.get(col_name)
    if spec is None:
        return None, None

    context = _neutral_row_context(row, config)
    if spec.when is not None and not spec.when(context):
        return None, None

    label = spec.label(context)
    url = spec.url(context)
    if not label or not url:
        return None, None
    return str(label), str(url)


def _column_width(col_name: str, df: pd.DataFrame, config: ReviewPipelineConfig) -> float:
    if col_name in config.width_overrides:
        return config.width_overrides[col_name]

    texts = [_cell_to_text(v) for v in df[col_name].tolist()]
    max_len = max([len(col_name)] + [len(t) for t in texts] + [0])

    if _starts_with_any(col_name, config.entity_prefixes) or _starts_with_any(
        col_name, config.business_prefixes
    ):
        return min(max(max_len * 0.9 + 2, 18), 40)
    if _starts_with_any(col_name, config.context_prefixes):
        return min(max(max_len * 0.9 + 2, 16), 32)
    if col_name in config.review_cols:
        return min(max(max_len * 1.1 + 2, 14), 28)
    if col_name in config.strat_cols:
        return min(max(max_len * 0.9 + 2, 12), 24)
    return min(max(max_len * 0.9 + 2, 12), 24)


def write_styled_excel(
    df: pd.DataFrame,
    out_path: str | Path,
    config: ReviewPipelineConfig,
    review_entity_count: int = 0,
) -> None:
    """Write one review workbook using the canonical EGE formatting."""
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    public = add_empty_review_columns(df, config)
    public = order_rows_for_review(public, config)

    # Hyperlink columns are generated by the Excel writer,
    # so they do not need to exist in the source dataframe.
    for col_name in config.hyperlinks:
        if col_name not in public.columns:
            public[col_name] = pd.NA

    public = public[ordered_columns(public, config)].copy()

    wb = Workbook()
    ws = wb.active
    ws.title = "Review"

    header_fill = PatternFill(fill_type="solid", fgColor=_to_argb(config.header_fill))
    header_font = Font(
        name="Calibri",
        size=11,
        bold=True,
        color=_to_argb(config.header_font_color),
    )
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    default_font = Font(name="Calibri", size=11, color="FF1F1F1F")
    review_font = Font(name="Calibri", size=11, bold=True, color="FF1F1F1F")
    hyperlink_font = Font(
        name="Calibri", size=11, color="FF366092", underline="single"
    )

    thin_side = Side(style="thin", color=_to_argb(config.thin_gray))
    medium_side = Side(style="medium", color=_to_argb(config.medium_gray))
    thin_border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side,
    )

    body_alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    center_alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)

    for col_idx, col_name in enumerate(public.columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = thin_border

    internal = _to_internal(public, config)
    entity_order = pd.unique(internal[_INTERNAL_ENTITY_ID])
    entity_to_band = {entity: idx for idx, entity in enumerate(entity_order)}
    boundary_entity = (
        entity_order[review_entity_count]
        if review_entity_count < len(entity_order)
        else None
    )
    seen_entities: set[Any] = set()

    for row_idx, (_, row) in enumerate(public.iterrows(), start=2):
        entity_value = row[config.entity_id_col]
        band_idx = entity_to_band.get(entity_value, 0)
        base_color = config.base_entity_colors[band_idx % len(config.base_entity_colors)]

        first_row_of_entity = entity_value not in seen_entities
        if first_row_of_entity:
            seen_entities.add(entity_value)

        top_side = (
            medium_side
            if (
                boundary_entity is not None
                and entity_value == boundary_entity
                and first_row_of_entity
            )
            else thin_side
        )

        for col_idx, col_name in enumerate(public.columns, start=1):
            value = row[col_name]
            link_label, url = _hyperlink_value(col_name, row, config)
            if col_name in config.hyperlinks:
                cell = ws.cell(
                    row=row_idx,
                    column=col_idx,
                    value=link_label if link_label else None,
                )
                if url:
                    cell.hyperlink = url
            else:
                cell = ws.cell(
                    row=row_idx,
                    column=col_idx,
                    value=None if pd.isna(value) else value,
                )

            cell.fill = PatternFill(
                fill_type="solid",
                fgColor=_to_argb(_column_fill(col_name, base_color, config)),
            )

            if col_name in config.review_cols:
                cell.font = review_font
                cell.border = Border(
                    left=medium_side,
                    right=medium_side,
                    top=top_side,
                    bottom=thin_side,
                )
            else:
                cell.font = hyperlink_font if url else default_font
                cell.border = Border(
                    left=thin_side,
                    right=thin_side,
                    top=top_side,
                    bottom=thin_side,
                )

            cell.alignment = (
                center_alignment if col_name in config.centered_cols else body_alignment
            )

            if col_name in config.date_cols and cell.value is not None:
                cell.number_format = config.date_number_format

    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    ws.row_dimensions[1].height = 26

    for col_idx, col_name in enumerate(public.columns, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = _column_width(
            col_name, public, config
        )

    wb.save(out_path)


# ---------------------------------------------------------------------------
# Export-directory handling
# ---------------------------------------------------------------------------
def _reset_output_directory(output_root: str | Path) -> Path:
    """Replace the review-export directory with an empty directory."""
    output_root = Path(output_root)

    if output_root.exists():
        shutil.rmtree(output_root)

    output_root.mkdir(parents=True, exist_ok=True)
    return output_root


# ---------------------------------------------------------------------------
# Public export API
# ---------------------------------------------------------------------------
def export_review_files(
    df: pd.DataFrame,
    strata_plan: pd.DataFrame,
    output_root: str | Path,
    config: ReviewPipelineConfig,
    display_total: int | None = None,
    confirmed_decisions: pd.DataFrame | None = None,
) -> dict[Any, dict[Any, pd.DataFrame]]:
    """Split, sample and export review files for either EJ/SIREN or EGE/SIRET.

    ``df`` is the proposals dataframe and does *not* need an assignment column.
    Assignment is looked up from ``strata_plan`` using the configured stratum
    columns.  If ``df`` already contains ``config.assignment_col``, that column
    is ignored and replaced by the strata-plan value.

    ``confirmed_decisions`` may contain one authoritative decision per entity.
    When provided, its configured review columns are copied to the first row of
    each matching sampled/displayed entity; sampling itself is unchanged.

    Returned dataframes always keep the configured/public source column names;
    ``entity_id`` and ``business_id`` exist only inside the engine.

    ``output_root`` represents one complete export batch. Its existing contents
    are removed before the new reviewer/stratum workbooks are written.
    """
    _validate_source_df(df, config)
    _, sample_map, _ = _normalized_strata_plan(strata_plan, config)
    df_with_assignment = _attach_assignment_from_strata_plan(
        df, strata_plan, config
    )

    display_total = config.display_total if display_total is None else display_total
    if display_total < 0:
        raise ValueError("display_total must be greater than or equal to zero.")

    # A new export is a complete batch: remove any workbooks left by a
    # previous export before writing the current reviewer/stratum files.
    output_root = _reset_output_directory(output_root)

    results: dict[Any, dict[Any, pd.DataFrame]] = {}

    for assignment, df_assignment in df_with_assignment.groupby(
        config.assignment_col,
        sort=False,
        dropna=False,
    ):
        assignment_dir = output_root / safe_name(assignment)
        assignment_dir.mkdir(parents=True, exist_ok=True)
        results[assignment] = {}

        for stratum_id, df_stratum in df_assignment.groupby(
            config.stratum_col,
            sort=False,
            dropna=False,
        ):
            requested_sample_size = int(
                sample_map.get(normalize_identifier(stratum_id), 0)
            )

            sampled, review_entity_count = sample_first_unique_entities_by_stratum(
                df_stratum,
                strata_plan=strata_plan,
                config=config,
                display_total=display_total,
            )
            sampled = apply_confirmed_decisions(
                sampled,
                confirmed_decisions=confirmed_decisions,
                config=config,
            )
            results[assignment][stratum_id] = sampled

            out_path = assignment_dir / (
                f"{safe_name(stratum_id)}__{requested_sample_size}.xlsx"
            )
            write_styled_excel(
                sampled,
                out_path,
                config=config,
                review_entity_count=review_entity_count,
            )

    return results


def summarize_review_export(
    results: Mapping[Any, Mapping[Any, pd.DataFrame]],
    strata_plan: pd.DataFrame,
    config: ReviewPipelineConfig,
    *,
    confirmed_decisions: pd.DataFrame | None = None,
) -> dict[str, int]:
    """Return compact counts describing one completed review-export batch."""
    _normalized_strata_plan(strata_plan, config)

    exported_entity_keys: set[str] = set()
    workbook_count = 0
    for strata in results.values():
        workbook_count += len(strata)
        for table in strata.values():
            if config.entity_id_col not in table.columns:
                raise ValueError(
                    f"Exported review table is missing {config.entity_id_col!r}."
                )
            exported_entity_keys.update(
                key
                for key in table[config.entity_id_col].map(normalize_identifier)
                if key
            )

    planned_review_entities = int(
        pd.to_numeric(
            strata_plan[config.sample_size_col],
            errors="raise",
        ).sum()
    )

    prefilled_entity_count = 0
    if confirmed_decisions is not None:
        if config.entity_id_col not in confirmed_decisions.columns:
            raise ValueError(
                "confirmed_decisions is missing the configured entity identifier "
                f"{config.entity_id_col!r}."
            )
        decision_keys = {
            key
            for key in confirmed_decisions[config.entity_id_col].map(
                normalize_identifier
            )
            if key
        }
        prefilled_entity_count = len(exported_entity_keys.intersection(decision_keys))

    return {
        "reviewer_count": len(results),
        "workbook_count": workbook_count,
        "planned_review_entity_count": planned_review_entities,
        "exported_entity_count": len(exported_entity_keys),
        "prefilled_entity_count": prefilled_entity_count,
    }


__all__ = [
    "HyperlinkSpec",
    "ReviewPipelineConfig",
    "add_empty_review_columns",
    "apply_confirmed_decisions",
    "export_review_files",
    "normalize_identifier",
    "order_rows_for_review",
    "ordered_columns",
    "safe_name",
    "sample_first_unique_entities_by_stratum",
    "summarize_review_export",
    "write_styled_excel",
]
